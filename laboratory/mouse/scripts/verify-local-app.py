from __future__ import annotations

import json
import io
import os
import sqlite3
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch

try:
    from fastapi.testclient import TestClient
except (ModuleNotFoundError, RuntimeError):
    TestClient = None

try:
    from openpyxl import Workbook, load_workbook
except ModuleNotFoundError:
    Workbook = None
    load_workbook = None

try:
    from PIL import Image, ImageDraw
except ModuleNotFoundError:
    Image = None
    ImageDraw = None


ROOT = Path(__file__).resolve().parents[1]
CLI_MAIN = ROOT / "mousedb" / "__main__.py"
CLI_PYTHON = ROOT / ".venv" / "Scripts" / "python.exe"


def assert_true(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def sample_blue_card_image_bytes() -> bytes:
    assert_true(Image is not None and ImageDraw is not None, "Pillow should be installed for ROI image verification.")
    image = Image.new("RGB", (960, 620), "#36aee2")
    draw = ImageDraw.Draw(image)
    line_color = "#1f4d69"
    for y in [60, 145, 230, 315, 390, 560]:
        draw.line((28, y, 930, y), fill=line_color, width=4)
    for x in [150, 360, 520, 760]:
        draw.line((x, 60, x, 390), fill=line_color, width=4)
    draw.text((45, 82), "Strain ApoM Tg/Tg", fill="#1c2678")
    draw.text((45, 162), "Sex F6", fill="#1c2678")
    draw.text((385, 162), "D.O.B 26.3.1-4", fill="#1c2678")
    draw.text((45, 250), "I.D Atg", fill="#1c2678")
    draw.text((650, 330), "LMO Y/N", fill="#1c2678")
    draw.text((45, 420), "Note", fill="#1c2678")
    draw.text((80, 470), "1 2 3 4 5", fill="#1c2678")
    output = io.BytesIO()
    image.save(output, format="JPEG", quality=92)
    return output.getvalue()


def run_cli(data_dir: Path, *args: str, expect_code: int = 0) -> subprocess.CompletedProcess[str]:
    python_executable = str(CLI_PYTHON) if CLI_PYTHON.exists() else sys.executable
    result = subprocess.run(
        [python_executable, "-m", "mousedb", "--db", str(data_dir / "mousedb.sqlite"), *args],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    assert_true(
        result.returncode == expect_code,
        f"CLI {' '.join(args)} returned {result.returncode}: stdout={result.stdout!r} stderr={result.stderr!r}",
    )
    return result


def main() -> None:
    for path in [
        ROOT / "app" / "main.py",
        ROOT / "app" / "db.py",
        ROOT / "app" / "storage.py",
        CLI_MAIN,
        ROOT / "static" / "index.html",
        ROOT / "requirements.txt",
        ROOT / "start.bat",
    ]:
        assert_true(path.exists(), f"Missing required local app file: {path}")

    fixture = json.loads((ROOT / "fixtures" / "sample_parse_results.json").read_text(encoding="utf-8"))
    assert_true(fixture.get("layer") == "parsed or intermediate result", "Fixture must stay non-canonical.")
    assert_true(len(fixture.get("records", [])) >= 3, "Fixture should contain parse records.")
    assert_true(Workbook is not None and load_workbook is not None, "openpyxl is required for workbook parsing and validation.")

    with tempfile.TemporaryDirectory() as source_dir:
        animal_path = Path(source_dir) / "legacy_animal.xlsx"
        animal_workbook = Workbook()
        animal_sheet = animal_workbook.active
        animal_sheet.title = "animal sheet"
        animal_sheet.append(["Cage No.", "Strain", "Sex", "I.D", "genotype", "DOB", "Mating date", "Pubs"])
        animal_sheet.append(["1", "ApoM Tg/Tg", "M", "MT321", "Tg/Tg", "2026-01-01", "2026-05-01", ""])
        animal_sheet.append(["", "", "F1", "9p", "pre_weaning", "2026-05-02", "", "2026-05-02 9p"])
        animal_workbook.save(animal_path)
        python_executable = str(CLI_PYTHON) if CLI_PYTHON.exists() else sys.executable
        animal_result = subprocess.run(
            [python_executable, str(ROOT / "scripts" / "parse_legacy_workbooks.py"), str(animal_path), "--kind", "animal"],
            cwd=ROOT,
            text=True,
            capture_output=True,
            check=False,
        )
        assert_true(animal_result.returncode == 0, f"Legacy animal parser failed: {animal_result.stderr}")
        animal_payload = json.loads(animal_result.stdout)
        assert_true(animal_payload["layer"] == "parsed or intermediate result", "Legacy workbook parser should stay non-canonical.")
        assert_true(animal_payload["source_layer"] == "export or view", "Legacy workbook source should be classified as export/view.")
        assert_true(animal_payload["rows"][0]["row_type"] == "parent_or_mouse_snapshot", "Animal parser should classify M rows as mouse snapshots.")
        assert_true(animal_payload["rows"][1]["row_type"] == "litter_or_offspring_snapshot", "Animal parser should classify F1 rows as litter snapshots.")
        assert_true(animal_payload["rows"][0]["source_cells"]["display_id"] == "D2", "Animal parser should preserve cell traceability.")

        separation_path = Path(source_dir) / "legacy_separation.xlsx"
        separation_workbook = Workbook()
        separation_sheet = separation_workbook.active
        separation_sheet.title = "separation"
        separation_sheet.append(["Strain", "Genotype", "total", "DOB", "WT", "Tg", "Sampling point"])
        separation_sheet.append(["ApoM Tg/Tg", "Tg/Tg", "M 3p", "2026-01-01", "", "3", "tail"])
        separation_workbook.save(separation_path)
        separation_result = subprocess.run(
            [python_executable, str(ROOT / "scripts" / "parse_legacy_workbooks.py"), str(separation_path), "--kind", "separation"],
            cwd=ROOT,
            text=True,
            capture_output=True,
            check=False,
        )
        assert_true(separation_result.returncode == 0, f"Legacy separation parser failed: {separation_result.stderr}")
        separation_payload = json.loads(separation_result.stdout)
        assert_true(separation_payload["workbook_kind"] == "legacy_separation_status", "Separation parser should label workbook kind.")
        assert_true(separation_payload["rows"][0]["sex_candidate"] == "male", "Separation parser should infer ASCII M sex labels.")
        assert_true(separation_payload["rows"][0]["count_candidate"] == 3, "Separation parser should infer pup counts.")

    sys.path.insert(0, str(ROOT))
    from app import db
    from scripts.parse_legacy_workbooks import parse_workbook

    with tempfile.TemporaryDirectory() as workbook_dir:
        workbook_root = Path(workbook_dir)
        animal_path = workbook_root / "legacy_animal.xlsx"
        animal_wb = Workbook()
        animal_ws = animal_wb.active
        animal_ws.title = "ApoM TgTg"
        animal_ws.append(["Cage No.", "Strain", "Sex", "I.D", "genotype", "DOB", "Mating date", "Pubs"])
        animal_ws.append(["1", "ApoM Tg/Tg", "\u2642", "MT 318 R'", "Tg", "25.10.20-28", "26.01.30", ""])
        animal_ws.append(["", "", "F1", "7p", "separated", "26.02.04", "", ""])
        animal_wb.save(animal_path)
        animal_payload = parse_workbook(animal_path, kind="animal")
        assert_true(animal_payload["layer"] == "parsed or intermediate result", "Legacy animal parser must stay non-canonical.")
        assert_true(animal_payload["source_layer"] == "export or view", "Legacy animal workbook must be classified as a view.")
        assert_true(len(animal_payload["rows"]) == 2, "Legacy animal parser row count is wrong.")
        assert_true(animal_payload["rows"][0]["source_cells"]["display_id"] == "D2", "Legacy animal parser must preserve cell traceability.")

        separation_path = workbook_root / "legacy_separation.xlsx"
        separation_wb = Workbook()
        separation_ws = separation_wb.active
        separation_ws.title = "ApoM TgTg"
        separation_ws.append(["Strain", "Genotype", "total", "DOB", "Genotype", "", "", "Sampling point"])
        separation_ws.append(["", "", "", "", "WT", "Tg", "", "10mths note"])
        separation_ws.append(["Apom Tg/Tg", "Apom Tg/Tg", "\u2642 2p", "25.05.07", "", "2", "", ""])
        separation_ws.append(["", "Apom Tg/Tg", "\u2640 6p", "26.02.18-24", "", "6", "", ""])
        separation_wb.save(separation_path)
        separation_payload = parse_workbook(separation_path, kind="separation")
        assert_true(separation_payload["source_layer"] == "export or view", "Legacy separation workbook must be classified as a view.")
        assert_true(len(separation_payload["rows"]) == 2, "Legacy separation parser should skip subheader-only rows.")
        assert_true(separation_payload["rows"][0]["count_candidate"] == 2, "Legacy separation parser should extract total counts as candidates.")
        assert_true(separation_payload["rows"][1]["sex_candidate"] == "female", "Legacy separation parser should extract sex as a candidate.")

    app = None
    if TestClient is not None:
        from app.main import app

    with tempfile.TemporaryDirectory() as old_schema_dir:
        db.DATA_DIR = Path(old_schema_dir)
        db.DB_PATH = Path(old_schema_dir) / "mouse_lims.sqlite"
        legacy_conn = sqlite3.connect(db.DB_PATH)
        try:
            legacy_conn.executescript(
                """
                CREATE TABLE mouse_master (
                    mouse_id TEXT PRIMARY KEY,
                    display_id TEXT NOT NULL
                );
                INSERT INTO mouse_master (mouse_id, display_id) VALUES ('legacy_mouse_001', 'LM001');
                CREATE TABLE card_note_item_log (
                    note_item_id TEXT PRIMARY KEY,
                    raw_line_text TEXT NOT NULL
                );
                CREATE TABLE genotyping_record (
                    genotyping_id TEXT PRIMARY KEY,
                    mouse_id TEXT,
                    sample_id TEXT,
                    sample_date TEXT,
                    result_date TEXT,
                    created_at TEXT
                );
                """
            )
            legacy_conn.commit()
        finally:
            legacy_conn.close()
        db.init_db()
        migrated_conn = sqlite3.connect(db.DB_PATH)
        try:
            mouse_columns = {
                row[1]
                for row in migrated_conn.execute("PRAGMA table_info(mouse_master)").fetchall()
            }
            note_columns = {
                row[1]
                for row in migrated_conn.execute("PRAGMA table_info(card_note_item_log)").fetchall()
            }
            genotype_columns = {
                row[1]
                for row in migrated_conn.execute("PRAGMA table_info(genotyping_record)").fetchall()
            }
            legacy_columns = {
                row[1]
                for row in migrated_conn.execute("PRAGMA table_info(legacy_workbook_row)").fetchall()
            }
        finally:
            migrated_conn.close()
        assert_true(
            {
                "target_match_status",
                "use_category",
                "next_action",
                "sample_id",
                "father_id",
                "mother_id",
                "litter_id",
                "source_record_id",
                "last_verified_at",
            }.issubset(mouse_columns),
            "Existing mouse_master tables should migrate to the genotyping, lineage, traceability, and verification schema.",
        )
        assert_true(
            {"parsed_ear_label_code", "strike_status", "needs_review"}.issubset(note_columns),
            "Existing card_note_item_log tables should migrate to the note parsing schema.",
        )
        assert_true(
            {"target_name", "normalized_result", "result_status", "updated_at"}.issubset(genotype_columns),
            "Existing genotyping_record tables should migrate to the result tracking schema.",
        )
        assert_true(
            "review_id" in legacy_columns,
            "Existing legacy_workbook_row tables should migrate to review linkage.",
        )

    if CLI_PYTHON.exists():
        with tempfile.TemporaryDirectory() as cli_dir:
            cli_data_dir = Path(cli_dir)
            initialized = json.loads(run_cli(cli_data_dir, "init", "--json").stdout)
            assert_true(initialized["initialized"] is True, "MouseDB CLI init should initialize the local database.")
            strain = json.loads(
                run_cli(
                    cli_data_dir,
                    "strain",
                    "add",
                    "--name",
                    "ApoM Tg/Tg",
                    "--source",
                    "manual",
                    "--json",
                ).stdout
            )
            assert_true(strain["strain_id"].startswith("STR-"), "MouseDB CLI strain add should create an external strain ID.")
            cage = json.loads(
                run_cli(
                    cli_data_dir,
                    "cage",
                    "add",
                    "--label",
                    "C-014",
                    "--type",
                    "holding",
                    "--json",
                ).stdout
            )
            assert_true(cage["cage_id"] == "C-014", "MouseDB CLI cage add should normalize cage IDs.")
            mouse = json.loads(
                run_cli(
                    cli_data_dir,
                    "mouse",
                    "add",
                    "--display-id",
                    "MT321",
                    "--strain",
                    strain["strain_id"],
                    "--sex",
                    "F",
                    "--dob",
                    "2025-10-20",
                    "--cage",
                    cage["cage_id"],
                    "--json",
                ).stdout
            )
            assert_true(mouse["display_id"] == "MT321", "MouseDB CLI mouse add should preserve display ID.")
            mate = json.loads(
                run_cli(
                    cli_data_dir,
                    "mouse",
                    "add",
                    "--display-id",
                    "MT322",
                    "--strain",
                    strain["strain_id"],
                    "--sex",
                    "M",
                    "--dob",
                    "2025-10-20",
                    "--cage",
                    cage["cage_id"],
                    "--json",
                ).stdout
            )
            assert_true(mate["mouse_id"] != mouse["mouse_id"], "MouseDB CLI should create distinct mouse IDs.")
            genotype = json.loads(
                run_cli(
                    cli_data_dir,
                    "genotype",
                    "record",
                    "--mouse",
                    mouse["mouse_id"],
                    "--result",
                    "Tg/Tg",
                    "--sample-id",
                    "S-MT321",
                    "--json",
                ).stdout
            )
            assert_true(genotype["result"] == "Tg/Tg", "MouseDB CLI genotype record should preserve result text.")
            mating = json.loads(
                run_cli(
                    cli_data_dir,
                    "mating",
                    "create",
                    "--male",
                    mate["mouse_id"],
                    "--female",
                    mouse["mouse_id"],
                    "--goal",
                    strain["strain_name"],
                    "--expected-genotype",
                    "Tg/Tg",
                    "--json",
                ).stdout
            )
            assert_true(mating["status"] == "active", "MouseDB CLI mating create should create an active mating.")
            litter = json.loads(
                run_cli(
                    cli_data_dir,
                    "litter",
                    "create",
                    "--mating",
                    mating["mating_id"],
                    "--number-born",
                    "6",
                    "--json",
                ).stdout
            )
            assert_true(litter["number_born"] == 6, "MouseDB CLI litter create should preserve litter counts.")
            summary = json.loads(run_cli(cli_data_dir, "colony", "summary", "--json").stdout)
            assert_true(summary["total_alive_mice"] >= 2, "MouseDB CLI colony summary should include live mouse totals.")
            assert_true(summary["active_matings"] == 1, "MouseDB CLI colony summary should include active mating counts.")

    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
        db.DATA_DIR = Path(temp_dir)
        db.DB_PATH = Path(temp_dir) / "mouse_lims.sqlite"
        db.init_db()
        conn = sqlite3.connect(db.DB_PATH)
        try:
            tables = {
                row[0]
                for row in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'")
            }
        finally:
            conn.close()
        assert_true(
            {
                "upload_batch",
                "photo_log",
                "parse_result",
                "review_queue",
                "action_log",
                "source_record",
                "strain_registry",
                "correction_log",
                "canonical_candidate",
                "export_log",
                "mouse_event",
                "genotyping_record",
                "strain_target_genotype",
                "genotype_status_master",
                "review_role_master",
                "review_priority_master",
                "cage_registry",
                "mouse_cage_assignment",
                "mating_registry",
                "mating_mouse",
                "litter_registry",
                "my_assigned_strain",
                "distribution_import",
                "distribution_assignment_row",
                "legacy_workbook_import",
                "legacy_workbook_row",
                "ear_label_master",
                "ear_label_alias",
                "mouse_master",
                "card_note_item_log",
            }.issubset(tables),
            "Local SQLite schema is incomplete.",
        )
        conn = sqlite3.connect(db.DB_PATH)
        try:
            master_rows = dict(
                conn.execute("SELECT ear_label_code, display_text FROM ear_label_master").fetchall()
            )
            ambiguous_alias = conn.execute(
                """
                SELECT confirmed
                FROM ear_label_alias
                WHERE raw_text = ? AND ear_label_code = ?
                """,
                ("R0", "R_CIRCLE"),
            ).fetchone()
            conn.execute(
                """
                INSERT INTO card_note_item_log
                    (note_item_id, line_number, raw_line_text, parsed_type, strike_status,
                     interpreted_status, parsed_mouse_display_id, parsed_ear_label_raw,
                     parsed_ear_label_code, confidence)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    "note_test_319",
                    1,
                    "319 L'",
                    "mouse_item",
                    "none",
                    "active",
                    "319",
                    "L'",
                    "L_PRIME",
                    0.98,
                ),
            )
            conn.execute(
                """
                INSERT INTO mouse_master
                    (mouse_id, display_id, raw_strain_text, dob_raw, dob_start, dob_end,
                     ear_label_raw, ear_label_code, source_note_item_id)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    "mouse_test_319_a",
                    "319",
                    "ApoM Tg/Tg",
                    "25.10.20-28",
                    "2025-10-20",
                    "2025-10-28",
                    "L'",
                    "L_PRIME",
                    "note_test_319",
                ),
            )
            conn.execute(
                """
                INSERT INTO mouse_master
                    (mouse_id, display_id, raw_strain_text, dob_start, ear_label_code)
                VALUES (?, ?, ?, ?, ?)
                """,
                ("mouse_test_319_b", "319", "Different strain candidate", "2026-01-01", "R_PRIME"),
            )
            mouse_defaults = conn.execute(
                """
                SELECT genotyping_status, next_action, status
                FROM mouse_master
                WHERE mouse_id = ?
                """,
                ("mouse_test_319_a",),
            ).fetchone()
            review_columns = {
                row[1]
                for row in conn.execute("PRAGMA table_info(review_queue)").fetchall()
            }
            action_columns = {
                row[1]
                for row in conn.execute("PRAGMA table_info(action_log)").fetchall()
            }
            export_columns = {
                row[1]
                for row in conn.execute("PRAGMA table_info(export_log)").fetchall()
            }
            photo_columns = {
                row[1]
                for row in conn.execute("PRAGMA table_info(photo_log)").fetchall()
            }
            same_display_count = conn.execute(
                "SELECT COUNT(*) FROM mouse_master WHERE display_id = ?",
                ("319",),
            ).fetchone()[0]
            genotype_statuses = {
                row[0]: row[1:]
                for row in conn.execute(
                    """
                    SELECT status_key, blocks_experiment, export_warning, legacy_genotyping_status
                    FROM genotype_status_master
                    """
                ).fetchall()
            }
            review_roles = {
                row[0]: row[1:]
                for row in conn.execute(
                    """
                    SELECT display_label, default_priority, active
                    FROM review_role_master
                    """
                ).fetchall()
            }
            review_priorities = {
                row[0]: row[1:]
                for row in conn.execute(
                    """
                    SELECT priority_key, severity_rank, export_blocking_hint
                    FROM review_priority_master
                    """
                ).fetchall()
            }
            conn.commit()
        finally:
            conn.close()
        assert_true(master_rows.get("R_CIRCLE") == "R\u00b0", "R_CIRCLE must use degree-sign display text.")
        assert_true(master_rows.get("L_CIRCLE") == "L\u00b0", "L_CIRCLE must use degree-sign display text.")
        assert_true(ambiguous_alias is not None, "Ambiguous R0 alias should be seeded for review.")
        assert_true(ambiguous_alias[0] == 0, "Ambiguous R0 alias must not be auto-confirmed.")
        assert_true(
            tuple(mouse_defaults) == ("not_sampled", "sample_needed", "active"),
            "Mouse workflow defaults are wrong.",
        )
        assert_true(
            {"assigned_role", "assigned_to", "priority"}.issubset(review_columns),
            "Review Queue should carry persona assignment and priority fields.",
        )
        assert_true(
            {"performed_by", "performed_role"}.issubset(action_columns),
            "Action Log should preserve who/which persona performed workflow changes.",
        )
        assert_true(
            {"generated_by", "generated_role"}.issubset(export_columns),
            "Export Log should preserve who/which persona generated export views.",
        )
        assert_true(
            "upload_batch_id" in photo_columns,
            "Photo log should link raw source photos to an upload batch.",
        )
        assert_true(same_display_count == 2, "Mouse display IDs must remain non-unique identity candidates.")
        assert_true(
            "confirmed" in genotype_statuses
            and genotype_statuses["confirmed"][0] == 0
            and genotype_statuses["confirmed"][2] == "resulted",
            "Genotype status vocabulary should mark confirmed results as resulted and not experiment-blocking.",
        )
        assert_true(
            "inconclusive" in genotype_statuses
            and genotype_statuses["inconclusive"][0] == 1
            and genotype_statuses["inconclusive"][1] == 1,
            "Genotype status vocabulary should keep inconclusive results blocking and export-warning.",
        )
        assert_true(
            "Experiment Planner" in review_roles
            and review_roles["Experiment Planner"][0] == "high"
            and review_roles["Experiment Planner"][1] == 1,
            "Review role master should seed experiment planning ownership without relying only on UI literals.",
        )
        assert_true(
            "high" in review_priorities
            and review_priorities["high"][0] == 1
            and review_priorities["high"][1] == 1,
            "Review priority master should preserve high priority as the top export-readiness hint.",
        )

        if TestClient is None:
            with db.connection() as conn:
                conn.execute(
                    """
                    INSERT INTO my_assigned_strain
                        (assigned_strain_id, display_name, aliases_json, source_type, assigned_at)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    ("assigned_strain_test", "ApoM Tg/Tg", '["ApoMtg/tg"]', "manual", "test"),
                )
                count = conn.execute("SELECT COUNT(*) AS count FROM my_assigned_strain").fetchone()["count"]
            assert_true(count == 1, "Assigned strain scope table did not accept a row.")
        else:
            with TestClient(app) as client:
                health = client.get("/api/health").json()
                assert_true(health["storage"] == "local-only", "Health endpoint should report local-only storage.")
                assert_true(
                    health["ai_draft"]["approval_required"] is True
                    and "selected photo" in health["ai_draft"]["payload_minimization"],
                    "Health endpoint should expose approval-gated AI draft status and payload minimization.",
                )
                with patch.dict(os.environ, {"OPENAI_API_KEY": ""}, clear=False):
                    set_ai_key = client.post("/api/ai-draft-settings", json={"api_key": "sk-test-session-key"})
                    assert_true(set_ai_key.status_code == 200, "Could not set temporary AI draft session key.")
                    set_ai_key_payload = set_ai_key.json()
                    assert_true(
                        set_ai_key_payload["ai_draft"]["available"] is True
                        and set_ai_key_payload["ai_draft"]["key_source"] == "session"
                        and "sk-test-session-key" not in json.dumps(set_ai_key_payload),
                        "AI draft session key should activate draft status without echoing the secret.",
                    )
                    cleared_ai_key = client.post("/api/ai-draft-settings", json={"api_key": ""})
                    assert_true(cleared_ai_key.status_code == 200, "Could not clear temporary AI draft session key.")
                    assert_true(
                        cleared_ai_key.json()["ai_draft"]["available"] is False,
                        "Clearing the session key should disable AI draft when no environment key is configured.",
                    )
                index_html = client.get("/").text
                assert_true("app-shell" in index_html and "Primary navigation" in index_html, "Local UI should use a persistent app shell with primary navigation.")
                assert_true('data-view-target="photo"' in index_html and "Photo Review Workbench" in index_html, "Photo Review Workbench should be the default operational view.")
                assert_true("reviewDetailPanel" in index_html and "inspect-review" in index_html, "Review Queue should expose a split list/detail workflow.")
                assert_true("Review Note Evidence" in index_html and "reviewEvidencePanel" in index_html, "Review detail should expose linked note-line and card snapshot evidence.")
                assert_true("Candidate Records" in index_html, "Local UI should expose candidate mouse records.")
                assert_true("Export Center" in index_html, "Local UI should frame workbook generation as an export center.")
                assert_true("Parsed Note Evidence" in index_html, "Local UI should expose parsed note evidence.")
                assert_true("Card Snapshots" in index_html and "cardSnapshotRows" in index_html, "Local UI should expose card transcription snapshots.")
                assert_true("Strain Registry" in index_html, "Local UI should expose strain registry.")
                assert_true(
                    "Linked gene/allele" in index_html and "strainAlleleSummary" in index_html,
                    "Local UI should surface normalized strain gene/allele links.",
                )
                assert_true("Source Evidence" in index_html, "Local UI should expose source evidence.")
                assert_true("Mouse Events" in index_html, "Local UI should expose mouse events.")
                assert_true("Correction Log" in index_html, "Local UI should expose correction history.")
                assert_true("Mouse Audit Trace" in index_html, "Local UI should expose per-mouse audit trace.")
                assert_true("audit-trail" in index_html, "Local UI should call the per-mouse audit trail API.")
                assert_true("Search & CSV Export" in index_html, "Local UI should expose search and CSV export.")
                assert_true("Download Genotyping Worklist" in index_html, "Local UI should expose genotyping worklist export.")
                assert_true("Download Ready CSV" in index_html, "Local UI should expose gated final CSV export.")
                assert_true("Cage View" in index_html, "Local UI should expose cage management.")
                assert_true("Breeding / Litter View" in index_html, "Local UI should expose mating and litter management.")
                assert_true("Create Offspring" in index_html, "Local UI should expose litter offspring generation.")
                assert_true("Complete Weaning" in index_html, "Local UI should expose litter weaning completion.")
                assert_true("Request Genotyping" in index_html, "Local UI should expose genotyping request workflow.")
                assert_true("requestLabelingRuleSet" in index_html, "Local UI should expose labeling rule selection for genotyping requests.")
                assert_true("/api/labeling-rule-sets" in index_html, "Local UI should fetch labeling rule policy.")
                assert_true("Target genotype" in index_html, "Local UI should expose configurable target genotype rules.")
                assert_true("genotypingDashboard" in index_html, "Local UI should expose genotyping dashboard cards.")
                assert_true("genotypeStatusRows" in index_html and "Genotype status vocabulary" in index_html, "Local UI should expose genotype status vocabulary.")
                assert_true("experimentReadinessRows" in index_html and "Experiment readiness" in index_html, "Local UI should expose experiment readiness planning.")
                assert_true("exportRows" in index_html, "Local UI should expose export preview rows.")
                assert_true("exportFilenames" in index_html, "Local UI should expose expected workbook filenames.")
                assert_true("exportReadinessCard" in index_html, "Local UI should expose an export readiness status card.")
                assert_true("exportExpectedFiles" in index_html, "Local UI should summarize expected final export files.")
                assert_true("exportBlockerList" in index_html, "Local UI should summarize open export blockers before table detail.")
                assert_true("exportBlockerRows" in index_html, "Local UI should expose export blockers.")
                assert_true("exportLogRows" in index_html, "Local UI should expose export history.")
                assert_true("Review Queue" in index_html and "Evidence" in index_html, "Local UI should show review evidence context.")
                assert_true("Extract & Save Review" in index_html, "Local UI should expose AI photo extraction as a saved review workflow.")
                assert_true("Load Audit Trace" in index_html, "Review detail should expose review audit trace loading.")
                assert_true("AI extraction skipped by user confirmation" in index_html, "AI photo extraction should stay user-confirmed before external inference.")
                assert_true("reviewStatusFilter" in index_html, "Local UI should expose review status filtering.")
                assert_true("reviewSeverityFilter" in index_html, "Local UI should expose review severity filtering.")
                assert_true("reviewEvidenceFilter" in index_html, "Local UI should expose review evidence filtering.")
                assert_true("reviewRoleFilter" in index_html and "reviewRoleRows" in index_html, "Review Queue should expose persona-based work queues from configured masters.")
                assert_true("reviewPriorityRows" in index_html and "/api/review-vocabulary" in index_html, "Review Queue should expose configured priority meanings.")
                assert_true("Last verified" in index_html and "last_verified_at" in index_html, "Mouse records should expose current-state verification timestamps.")
                assert_true("review-resolution-note" in index_html, "Local UI should resolve reviews inline instead of using prompt-only workflow.")
                assert_true("Mouse Audit Trace" in index_html, "Local UI should expose mouse audit trace view.")
                assert_true("auditTraceRows" in index_html, "Local UI should render audit trace rows.")
                assert_true("Deactivate" in index_html, "Local UI should expose assigned strain deactivation.")
                assert_true("Distribution Assignment Import" in index_html, "Local UI should expose distribution import.")
                assert_true("Legacy Workbook Import" in index_html, "Local UI should expose legacy workbook import.")
                assert_true("legacyWorkbookKind" in index_html, "Local UI should expose legacy workbook kind selection.")
                assert_true("legacyWorkbookRows" in index_html, "Local UI should render legacy workbook rows.")
                assert_true("Review items" in index_html, "Local UI should expose legacy workbook review item counts.")
                assert_true(
                    "Strain registry candidates" in index_html and "review gene/allele" in index_html,
                    "Local UI should surface legacy strain registry candidates without inferred gene/allele values.",
                )
                assert_true(
                    "apply_strain_registry_candidate" in index_html
                    and "reviewed-gene-symbol" in index_html
                    and "reviewed-allele-name" in index_html,
                    "Local UI should require reviewed strain/gene/allele values before applying legacy registry candidates.",
                )
                assert_true("reviewed-existing-strain-id" in index_html, "Local UI should support explicit existing-strain linkage.")
                assert_true("Legacy decision" in index_html, "Local UI should expose legacy review decision controls.")
                assert_true("Create Missing Photo Reviews" in index_html, "Local UI should expose photo review candidate creation.")
                assert_true("Evidence Reconciliation" in index_html, "Local UI should expose photo/workbook reconciliation.")
                assert_true("Evidence Comparison" in index_html, "Local UI should expose photo/workbook comparison.")
                assert_true("evidenceComparisonRows" in index_html, "Local UI should render evidence comparison rows.")
                assert_true("comparisonDetailPanel" in index_html and "inspect-comparison" in index_html, "Evidence Comparison should expose a decision table with a detail inspector.")
                assert_true("Matched fields" in index_html and "Mismatched fields" in index_html, "Evidence Comparison inspector should show matched and mismatched field context.")
                assert_true("comparisonReviewButton" in index_html, "Local UI should create comparison review candidates explicitly.")
                assert_true("Review state" in index_html, "Local UI should show whether comparison reviews are open, resolved, or not created.")
                assert_true("Canonical Candidate Drafts" in index_html, "Local UI should expose canonical candidate drafts.")
                assert_true("canonicalCandidateRows" in index_html, "Local UI should render canonical candidate drafts.")
                assert_true("Preview" in index_html and "apply-preview" in index_html, "Local UI should expose canonical apply preview.")
                assert_true(
                    "canonicalApplyPreviewPanel" in index_html
                    and "renderCanonicalApplyPreview" in index_html
                    and "Applying this candidate writes canonical mouse state" in index_html,
                    "Canonical apply preview should show proposed writes, blockers, and a clear canonical-write warning.",
                )
                assert_true("Void Applied" in index_html and "/audit" in index_html, "Local UI should expose applied candidate audit and void controls.")
                assert_true("transcriptionPhotoPreview" in index_html, "Local UI should preview selected raw cage-card photos.")
                assert_true("/api/photos/${encodeURIComponent(photo.photo_id)}/image" in index_html, "Local UI should load raw photo evidence from the local image endpoint.")
                assert_true("firstPendingPhotoButton" in index_html and "photoQueueSummary" in index_html, "Local UI should expose batch photo review queue controls.")
                assert_true("/api/photo-review-workbench" in index_html, "Local UI should load the batch photo review workbench view.")
                assert_true("aiDraftButton" in index_html, "Local UI should expose approval-gated AI transcription drafts.")
                assert_true("ai-transcription-draft" in index_html, "Local UI should call the AI draft endpoint only from the review form.")
                assert_true("openAiApiKeyInput" in index_html, "Local UI should allow a temporary session API key for AI drafts.")
                assert_true("ai-draft-settings" in index_html, "Local UI should update AI draft settings without storing the key in app records.")
                assert_true("transcriptionSexRaw" in index_html and "transcriptionIdRaw" in index_html, "Manual transcription should capture raw Sex and I.D card fields.")
                assert_true("inferCardTypeFromSexSelection" in index_html and "Mating" in index_html, "Mixed sex selection should default the parsed card type to Mating.")
                assert_true("unlabeled" in index_html or "1 2 3 4 5" in index_html, "Manual transcription should make numeric-only temporary labels visible in note entry.")
                assert_true("note-label-decision" in index_html, "Review UI should expose structured transcription label correction controls.")
                assert_true("ear-label-code" in index_html and "R_PRIME" in index_html, "Review UI should resolve bounded ear-label corrections with a select control.")
                assert_true(
                    "review-source-preview" in index_html
                    and "reviewSourceEvidencePanel" in index_html,
                    "Review detail should expose raw photo, note anchor, and card snapshot evidence in one panel.",
                )
                assert_true(
                    "Card Snapshots" in index_html
                    and "cardSnapshotRows" in index_html
                    and "/api/card-snapshots" in index_html,
                    "Local UI should expose parsed cage-card snapshots and render them from the card snapshot API.",
                )
                assert_true(
                    "photoZoomInButton" in index_html
                    and "photoRotateRightButton" in index_html
                    and "transcriptionPhotoEvidenceMeta" in index_html,
                    "Photo review should expose display-only zoom/rotation controls and source evidence metadata.",
                )
                assert_true(
                    "roiPreviewPanel" in index_html
                    and "/roi-preview" in index_html
                    and "renderRoiPreview" in index_html,
                    "Photo review should expose ROI card/field crop previews for selected cage-card photos.",
                )
                assert_true(
                    "extractionEvidencePanel" in index_html
                    and "uncertainFieldPills" in index_html
                    and "plausibilityWarningPills" in index_html
                    and "symbolConfusionPills" in index_html
                    and "renderExtractionEvidence" in index_html,
                    "Photo review should expose reviewable AI uncertainty, plausibility warnings, symbol confusions, and raw visible text evidence.",
                )
                assert_true(
                    "focus-uncertain-field" in index_html
                    and "needs-review" in index_html
                    and "focusUncertainField" in index_html,
                    "Photo review should make uncertain extracted fields actionable from the evidence panel.",
                )
                assert_true("multiple" in index_html and "Upload Photos" in index_html, "Local UI should support multi-photo upload.")
                assert_true("uploadBatchRows" in index_html and "/api/upload-batches" in index_html, "Local UI should expose upload batch tracking.")
                assert_true("uploadBatchReleasePanel" in index_html and "Preview release" in index_html, "Local UI should expose upload batch release checks.")
                assert_true("Photo worklist" in index_html and "next_action" in index_html, "Local UI should show photo-level batch release blockers.")
                assert_true("open-batch-worklist-target" in index_html and "item.action_target_type" in index_html, "Local UI should link batch release worklist rows to the next review action.")
                assert_true("manual_parse_id" in index_html and "scoped comparison review" in index_html, "Batch release action links should create scoped comparison review items.")
                assert_true("create-batch-comparison-reviews" in index_html and "upload_batch_id" in index_html, "Batch release preview should create comparison reviews scoped to the selected upload batch.")
                assert_true("/release-preview" in index_html and "Close Batch" in index_html, "Local UI should preview and close upload batches only after release checks.")
                assert_true("Manual Photo Transcription" in index_html, "Local UI should expose manual photo transcription.")
                assert_true("Colony Dashboard" in index_html, "Local UI should expose the colony visualization dashboard.")
                assert_true("Mouse Detail" in index_html, "Local UI should expose the mouse detail visualization.")
                assert_true("Mouse Audit Trace" in index_html, "Local UI should expose mouse audit trace.")
                assert_true("auditTraceRows" in index_html, "Local UI should render audit trace rows.")
                assert_true("Strain Detail" in index_html, "Local UI should expose the strain detail visualization.")
                assert_true("Evidence & Review Readiness" in index_html, "Local UI should expose visualization evidence readiness.")
                assert_true("renderVisualizations" in index_html, "Local UI should render visualizations from API data.")
                assert_true("vizHeatmapHead" in index_html, "Genotype heatmap should be driven by rendered data labels.")
                assert_true("vizQualityRows" in index_html, "Visualization readiness should render data quality rows.")
                assert_true("demo-note-1" not in index_html, "Mouse detail visualization should not use demo source evidence.")
                assert_true("selectedStrainMice.length || aliveMice" not in index_html, "Strain detail active mice should not fall back to colony-wide counts.")
                assert_true("EXP-2026-041" not in index_html, "Strain visualization should not hard-code experiment IDs.")
                assert_true(r"\p{L}\p{N}" in index_html, "Local UI strain matching key should preserve Korean and other letter/number strain text.")
                assert_true(client.get("/api/assigned-strains").json() == [], "Assigned strain scope should start empty.")
                assert_true(client.get("/api/source-records").json() == [], "Source evidence should start empty.")
                assert_true(client.get("/api/strains").json() == [], "Strain registry should start empty.")
                genotype_vocabulary = client.get("/api/genotype-status-vocabulary").json()
                assert_true(
                    any(item["status_key"] == "confirmed" and item["blocks_experiment"] is False for item in genotype_vocabulary)
                    and any(item["status_key"] == "pending" and item["export_warning"] is True for item in genotype_vocabulary),
                    "Genotype status vocabulary API should distinguish confirmed from pending/review states.",
                )
                review_vocabulary = client.get("/api/review-vocabulary").json()
                assert_true(
                    review_vocabulary["boundary"] == "canonical structured state"
                    and any(item["display_label"] == "Colony Reviewer" for item in review_vocabulary["roles"])
                    and any(item["priority_key"] == "high" and item["export_blocking_hint"] is True for item in review_vocabulary["priorities"]),
                    "Review vocabulary API should expose configured personas and priority meanings.",
                )
                initial_experiment_readiness = client.get("/api/experiment-readiness").json()
                assert_true(
                    initial_experiment_readiness["boundary"] == "export or view"
                    and "rows" in initial_experiment_readiness
                    and initial_experiment_readiness["total_count"] >= 0,
                    "Experiment readiness should be an export/view planning surface.",
                )
                strain = client.post(
                    "/api/strains",
                    json={
                        "strain_name": "PV-Cre",
                        "gene": "Pvalb",
                        "allele": "Pvalb-IRES-Cre",
                        "background": "C57BL/6J",
                        "source": "manual",
                        "status": "active",
                    },
                )
                assert_true(strain.status_code == 200, "Could not create strain registry entry.")
                strain_payload = strain.json()
                assert_true(strain_payload["source_record_id"], "Strain entry should create source evidence.")
                strains = client.get("/api/strains").json()
                assert_true(strains[0]["strain_name"] == "PV-Cre", "Strain registry did not persist entry.")
                genes = client.get("/api/genes").json()
                alleles = client.get("/api/alleles").json()
                assert_true(
                    any(item["gene_symbol"] == "Pvalb" for item in genes),
                    "Strain creation should populate the normalized gene registry.",
                )
                assert_true(
                    any(item["allele_name"] == "Pvalb-IRES-Cre" and item["gene_symbol"] == "Pvalb" for item in alleles),
                    "Strain creation should populate the normalized allele registry.",
                )
                assert_true(
                    strains[0]["alleles"]
                    and strains[0]["alleles"][0]["allele_name"] == "Pvalb-IRES-Cre"
                    and strains[0]["alleles"][0]["gene_symbol"] == "Pvalb",
                    "Strain list should expose normalized strain-allele links while preserving legacy text fields.",
                )
                gene_update = client.patch(
                    f"/api/genes/{genes[0]['gene_id']}",
                    json={"full_name": "Parvalbumin", "description": "Curated strain registry metadata."},
                )
                allele_update = client.patch(
                    f"/api/alleles/{alleles[0]['allele_id']}",
                    json={"description": "IRES-Cre driver allele", "allele_type": "transgene"},
                )
                assert_true(gene_update.status_code == 200 and allele_update.status_code == 200, "Could not update gene/allele metadata.")
                updated_gene = client.get("/api/genes").json()[0]
                updated_allele = client.get("/api/alleles").json()[0]
                assert_true(
                    updated_gene["gene_symbol"] == "Pvalb"
                    and updated_gene["full_name"] == "Parvalbumin"
                    and updated_gene["description"] == "Curated strain registry metadata.",
                    "Gene metadata updates should preserve raw symbol while exposing curated text.",
                )
                assert_true(
                    updated_allele["allele_name"] == "Pvalb-IRES-Cre"
                    and updated_allele["description"] == "IRES-Cre driver allele"
                    and updated_allele["allele_type"] == "transgene",
                    "Allele metadata updates should preserve raw allele symbol while exposing curated text.",
                )
                source_records = client.get("/api/source-records").json()
                assert_true(
                    any(item["source_type"] == "manual_entry" for item in source_records),
                    "Manual strain creation should leave source evidence.",
                )
                distribution = client.post(
                    "/api/distribution-imports",
                    json={
                        "layer": "parsed or intermediate result",
                        "source_file_name": "distribution_test.xlsx",
                        "sheet_name": "Mating",
                        "rows": [
                            {
                                "institution_or_group": "Vet Med",
                                "responsible_person_raw": "Jang S.",
                                "mating_type_raw": "ApoMtg/tg",
                                "cage_count_raw": "6",
                                "source_row_number": 35,
                                "source_cells": {"mating_type": "B35"},
                                "review_status": "candidate",
                            }
                        ],
                    },
                )
                assert_true(distribution.status_code == 200, "Could not store distribution import JSON.")
                distribution_payload = distribution.json()
                assert_true(distribution_payload["stored_rows"] == 1, "Distribution import row count is wrong.")
                assert_true(distribution_payload["source_record_id"], "Distribution import should create source evidence.")
                distribution_imports = client.get("/api/distribution-imports").json()
                stored_distribution = next(
                    (
                        item
                        for item in distribution_imports
                        if item["distribution_import_id"] == distribution_payload["distribution_import_id"]
                    ),
                    None,
                )
                assert_true(stored_distribution is not None, "Distribution import list should include the stored import.")
                assert_true(
                    stored_distribution["rows"][0]["traceability"]["source_cells"]["mating_type"] == "B35",
                    "Distribution import should preserve source cell traceability.",
                )
                mice_before_legacy = client.get("/api/mice").json()
                with tempfile.TemporaryDirectory() as legacy_upload_dir:
                    legacy_path = Path(legacy_upload_dir) / "legacy_animal_upload.xlsx"
                    legacy_wb = Workbook()
                    legacy_ws = legacy_wb.active
                    legacy_ws.title = "animal sheet"
                    legacy_ws.append(["Cage No.", "Strain", "Sex", "I.D", "genotype", "DOB", "Mating date", "Pubs"])
                    legacy_ws.append(["C-014", "ApoM Tg/Tg", "M", "MT321", "Tg/Tg", "2026-01-01", "2026-05-01", ""])
                    legacy_ws.append(["", "", "F1", "9p", "pre_weaning", "2026-05-02", "", "2026-05-02 9p"])
                    legacy_wb.save(legacy_path)
                    with legacy_path.open("rb") as legacy_file:
                        legacy_import = client.post(
                            "/api/legacy-workbook-imports",
                            data={"kind": "animal"},
                            files={
                                "file": (
                                    "legacy_animal_upload.xlsx",
                                    legacy_file,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                )
                            },
                        )
                assert_true(legacy_import.status_code == 200, f"Could not import legacy workbook: {legacy_import.text}")
                legacy_payload = legacy_import.json()
                assert_true(legacy_payload["boundary"] == "parsed or intermediate result", "Legacy workbook import should stay non-canonical.")
                assert_true(legacy_payload["stored_rows"] == 2, "Legacy workbook import row count is wrong.")
                assert_true(legacy_payload["parse_id"], "Legacy workbook import should create parse evidence for review items.")
                assert_true(
                    legacy_payload["created_review_items"] == 3
                    and legacy_payload["strain_registry_candidate_count"] == 1,
                    "Legacy workbook import should open row reviews plus strain registry candidate review.",
                )
                assert_true(legacy_payload["source_record_id"], "Legacy workbook import should create source evidence.")
                legacy_imports = client.get("/api/legacy-workbook-imports").json()
                stored_legacy = next(
                    (
                        item
                        for item in legacy_imports
                        if item["legacy_import_id"] == legacy_payload["legacy_import_id"]
                    ),
                    None,
                )
                assert_true(stored_legacy is not None, "Legacy workbook import list should include the stored import.")
                assert_true(stored_legacy["workbook_kind"] == "legacy_animal_sheet", "Legacy workbook kind was not preserved.")
                assert_true(stored_legacy["review_count"] == 3, "Legacy workbook list should expose linked review items.")
                assert_true(stored_legacy["open_review_count"] == 3, "Legacy workbook list should expose open review items.")
                assert_true(
                    stored_legacy["strain_registry_candidates"][0]["normalized_candidate"]["gene_symbol"] == "",
                    "Legacy workbook strain registry candidate should not infer gene symbols.",
                )
                assert_true(
                    stored_legacy["rows"][0]["raw_row"]["source_cells"]["display_id"] == "D2",
                    "Legacy workbook import should preserve source cell traceability.",
                )
                legacy_review_items = [
                    item
                    for item in client.get("/api/review-items").json()
                    if item["parse_id"] == legacy_payload["parse_id"]
                ]
                assert_true(len(legacy_review_items) == 3, "Legacy workbook row and registry reviews should be visible in Review Queue.")
                assert_true(
                    all(item["status"] == "open" for item in legacy_review_items),
                    "Legacy workbook review items should start open.",
                )
                legacy_row_review_items = [
                    item
                    for item in legacy_review_items
                    if item["issue"] == "Legacy workbook row requires review"
                ]
                registry_review_items = [
                    item
                    for item in legacy_review_items
                    if item["issue"] == "Legacy strain registry candidate requires review"
                ]
                assert_true(len(legacy_row_review_items) == 2, "Legacy workbook rows should still get row-level reviews.")
                assert_true(
                    len(registry_review_items) == 1
                    and registry_review_items[0]["assigned_role"] == "Strain Curator",
                    "Legacy workbook strain registry candidates should route to Strain Curator review.",
                )
                assert_true(
                    any("display_id" in item["current_value"] and "D2" in item["review_reason"] for item in legacy_row_review_items),
                    "Legacy workbook review items should preserve row anchors and source cell evidence.",
                )
                rejected_legacy_review = legacy_row_review_items[0]
                legacy_reject = client.post(
                    f"/api/review-items/{rejected_legacy_review['review_id']}/resolve",
                    json={
                        "resolution_note": "Reject predecessor workbook row after latest photo comparison.",
                        "resolved_value": "not accepted",
                        "legacy_decision": "reject_legacy_candidate",
                    },
                )
                assert_true(legacy_reject.status_code == 200, "Could not reject a legacy workbook review row.")
                legacy_reject_payload = legacy_reject.json()
                assert_true(
                    legacy_reject_payload["legacy_row_review_status"] == "rejected",
                    "Legacy review rejection should mark the workbook row rejected.",
                )
                updated_legacy_imports = client.get("/api/legacy-workbook-imports").json()
                updated_stored_legacy = next(
                    item
                    for item in updated_legacy_imports
                    if item["legacy_import_id"] == legacy_payload["legacy_import_id"]
                )
                assert_true(
                    updated_stored_legacy["open_review_count"] == 2,
                    "Resolving a legacy row should reduce the legacy import open review count.",
                )
                rejected_rows = [
                    row
                    for row in updated_stored_legacy["rows"]
                    if row["review_id"] == rejected_legacy_review["review_id"]
                ]
                assert_true(
                    rejected_rows and rejected_rows[0]["review_status"] == "rejected",
                    "Legacy workbook row should expose the reviewer decision status.",
                )
                with sqlite3.connect(db.DB_PATH) as conn:
                    conn.row_factory = sqlite3.Row
                    legacy_parse = conn.execute(
                        "SELECT parse_id, source_name, status, needs_review FROM parse_result WHERE parse_id = ?",
                        (legacy_payload["parse_id"],),
                    ).fetchone()
                assert_true(legacy_parse is not None, "Legacy workbook review parse_result should be persisted.")
                assert_true(legacy_parse["source_name"] == "legacy_animal_upload.xlsx", "Legacy parse should keep source filename.")
                assert_true(legacy_parse["status"] == "review", "Legacy parse should remain in review state.")
                assert_true(legacy_parse["needs_review"] == 1, "Legacy parse should require review.")
                assert_true(
                    client.get("/api/mice").json() == mice_before_legacy,
                    "Legacy workbook import should not write canonical mouse state.",
                )
                upload_batch = client.post(
                    "/api/upload-batches",
                    json={"batch_label": "Morning cage card set", "expected_photo_count": 2, "note": "Verification batch."},
                )
                assert_true(upload_batch.status_code == 200, f"Could not create upload batch: {upload_batch.text}")
                upload_batch_payload = upload_batch.json()
                assert_true(
                    upload_batch_payload["source_layer"] == "raw source"
                    and upload_batch_payload["expected_photo_count"] == 2,
                    "Upload batch creation should classify the batch as raw source evidence.",
                )
                photo_upload = client.post(
                    "/api/photos",
                    data={"upload_batch_id": upload_batch_payload["upload_batch_id"]},
                    files={"file": ("latest_card.jpg", io.BytesIO(b"fake local image bytes"), "image/jpeg")},
                )
                assert_true(photo_upload.status_code == 200, f"Could not upload source photo: {photo_upload.text}")
                photo_payload = photo_upload.json()
                assert_true(photo_payload["status"] == "review_pending", "Uploaded photos should enter manual review state.")
                assert_true(photo_payload["upload_batch_id"] == upload_batch_payload["upload_batch_id"], "Uploaded photo should preserve its upload batch link.")
                assert_true(photo_payload["review_candidate"]["review_id"], "Photo upload should create a review candidate.")
                photo_image = client.get(f"/api/photos/{photo_payload['photo_id']}/image")
                assert_true(photo_image.status_code == 200, "Uploaded photo raw evidence image should be retrievable locally.")
                assert_true(photo_image.content == b"fake local image bytes", "Photo image endpoint should return the preserved raw upload bytes.")
                assert_true(
                    photo_image.headers["content-type"].startswith("image/"),
                    "Photo image endpoint should preserve an image content type.",
                )
                photo_image.close()
                invalid_roi_preview = client.get(f"/api/photos/{photo_payload['photo_id']}/roi-preview")
                assert_true(
                    invalid_roi_preview.status_code == 415,
                    "ROI preview should reject unreadable uploaded image bytes without changing raw photo evidence.",
                )
                photo_workbench = client.get("/api/photo-review-workbench").json()
                assert_true(photo_workbench["boundary"] == "export or view", "Photo review workbench should remain a review/export view.")
                assert_true(
                    photo_workbench["batch_count"] >= 1
                    and any(batch["upload_batch_id"] == upload_batch_payload["upload_batch_id"] for batch in photo_workbench["batches"]),
                    "Photo review workbench should summarize upload batches.",
                )
                assert_true(photo_workbench["pending_transcription_count"] >= 1, "Untranscribed photos should be queued for manual transcription.")
                workbench_row = next(item for item in photo_workbench["rows"] if item["photo_id"] == photo_payload["photo_id"])
                assert_true(
                    workbench_row["next_action"] == "transcribe_photo"
                    and workbench_row["image_url"].endswith("/image")
                    and workbench_row["upload_batch_id"] == upload_batch_payload["upload_batch_id"],
                    "Photo review workbench should point raw photos to transcription and local image evidence.",
                )
                upload_batches = client.get("/api/upload-batches").json()
                selected_batch = next(item for item in upload_batches["rows"] if item["upload_batch_id"] == upload_batch_payload["upload_batch_id"])
                assert_true(
                    selected_batch["photo_count"] == 1
                    and selected_batch["pending_transcription_count"] == 1
                    and selected_batch["open_review_count"] >= 1,
                    "Upload batch summary should count photos, transcription pending state, and open reviews.",
                )
                batch_release_preview = client.get(
                    f"/api/upload-batches/{upload_batch_payload['upload_batch_id']}/release-preview"
                )
                assert_true(batch_release_preview.status_code == 200, "Upload batch release preview endpoint failed.")
                batch_release_payload = batch_release_preview.json()
                assert_true(
                    batch_release_payload["boundary"] == "export or view"
                    and batch_release_payload["ready"] is False
                    and len(batch_release_payload["worklist"]) == 1
                    and batch_release_payload["worklist"][0]["photo_id"] == photo_payload["photo_id"]
                    and batch_release_payload["worklist"][0]["next_action"] == "transcribe_photo"
                    and batch_release_payload["worklist"][0]["action_target_type"] == "photo"
                    and batch_release_payload["worklist"][0]["action_target_id"] == photo_payload["photo_id"]
                    and batch_release_payload["worklist"][0]["image_url"].endswith("/image")
                    and any(item["key"] == "transcription_complete" for item in batch_release_payload["blockers"])
                    and any(item["key"] == "canonical_mapping_applied" for item in batch_release_payload["blockers"])
                    and any(item["key"] == "photo_worklist_clear" for item in batch_release_payload["blockers"]),
                    "Upload batch release preview should identify unfinished transcription and missing canonical mapping at batch and photo level.",
                )
                blocked_batch_release = client.post(
                    f"/api/upload-batches/{upload_batch_payload['upload_batch_id']}/release"
                )
                assert_true(
                    blocked_batch_release.status_code == 409,
                    "Upload batch release should not close a batch with blockers.",
                )
                blocked_ai_draft = client.post(
                    f"/api/photos/{photo_payload['photo_id']}/ai-transcription-draft",
                    json={"approved_external_inference": False},
                )
                assert_true(
                    blocked_ai_draft.status_code == 403,
                    "AI draft transcription should require explicit per-request external inference approval.",
                )
                with patch.dict(os.environ, {"OPENAI_API_KEY": ""}, clear=False):
                    unavailable_ai_draft = client.post(
                        f"/api/photos/{photo_payload['photo_id']}/ai-transcription-draft",
                        json={"approved_external_inference": True},
                    )
                assert_true(
                    unavailable_ai_draft.status_code == 503,
                    "AI draft transcription should stay unavailable without an explicit API key.",
                )
                ai_photo_upload = client.post(
                    "/api/photos",
                    files={"file": ("ai_card.jpg", io.BytesIO(sample_blue_card_image_bytes()), "image/jpeg")},
                )
                assert_true(ai_photo_upload.status_code == 200, f"Could not upload AI extraction photo: {ai_photo_upload.text}")
                ai_photo_payload = ai_photo_upload.json()

                class FakeOpenAIResponse:
                    def raise_for_status(self) -> None:
                        return None

                    def json(self) -> dict[str, str]:
                        draft = {
                            "card_type": "Separated",
                            "raw_strain": "ApoM Tg/Tg",
                            "matched_strain": "ApoM Tg/Tg",
                            "sex_raw": "\u2640",
                            "id_raw": "MT",
                            "dob_raw": "25.10.20-28",
                            "dob_normalized": "",
                            "mating_date_raw": "",
                            "mating_date_normalized": "",
                            "lmo_raw": "",
                            "mouse_count": "2 total",
                            "notes": [
                                {"raw": "MT401 R'", "meaning": "mouse", "strike": "none", "confidence": 91},
                                {"raw": "MT402 L'", "meaning": "mouse", "strike": "none", "confidence": 90},
                            ],
                            "raw_visible_text_lines": ["ApoM Tg/Tg", "\u2640 2 total", "MT401 R'", "MT402 L'"],
                            "symbol_confusions": ["F vs \u2640 resolved from Sex ROI"],
                            "confidence": 88,
                            "uncertain_fields": ["dob_normalized"],
                            "reviewer_note": "Drafted from image.",
                        }
                        return {"output_text": json.dumps(draft)}

                class FakeOpenAIClient:
                    def __init__(self, *args, **kwargs) -> None:
                        self.request_json = None

                    def __enter__(self):
                        return self

                    def __exit__(self, *args) -> None:
                        return None

                    def post(self, url, headers=None, json=None):
                        self.request_json = json
                        assert_true(url.endswith("/v1/responses"), "AI extraction should call the Responses API.")
                        payload_text = json["input"][0]["content"][0]["text"]
                        assert_true(
                            "Assigned strain scope" in payload_text,
                            "AI extraction prompt should include only bounded assigned strain context.",
                        )
                        assert_true(
                            "Extraction image mode: roi_field_crops" in payload_text
                            and "target_fields" in payload_text,
                            "AI extraction prompt should describe fixed ROI regions and their target fields.",
                        )
                        image_parts = [
                            item for item in json["input"][0]["content"]
                            if item.get("type") == "input_image"
                        ]
                        assert_true(
                            len(image_parts) > 1,
                            "AI extraction should send normalized card and field ROI crops instead of only one full photo.",
                        )
                        label_text = "\n".join(
                            item.get("text", "")
                            for item in json["input"][0]["content"]
                            if item.get("type") == "input_text"
                        )
                        assert_true(
                            "ROI label=sex_raw" in label_text
                            and "ROI label=dob_raw" in label_text
                            and "ROI label=notes" in label_text,
                            "AI extraction should label field-specific ROI crops for the model.",
                        )
                        return FakeOpenAIResponse()

                with patch.dict(os.environ, {"OPENAI_API_KEY": "sk-test-ai-extract"}, clear=False):
                    with patch("app.main.httpx.Client", FakeOpenAIClient):
                        ai_extraction = client.post(
                            f"/api/photos/{ai_photo_payload['photo_id']}/ai-extract-transcription",
                            json={"approved_external_inference": True, "detail": "low"},
                        )
                assert_true(ai_extraction.status_code == 200, f"Could not save AI extraction: {ai_extraction.text}")
                ai_extraction_payload = ai_extraction.json()
                assert_true(
                    ai_extraction_payload["source_name"] == "ai_photo_extraction"
                    and ai_extraction_payload["created_note_items"] == 2
                    and ai_extraction_payload["created_mouse_candidates"] == 0,
                    "AI extraction should save parsed note evidence for review without canonical mouse writes.",
                )
                assert_true(
                    ai_extraction_payload["extraction_image_mode"] == "roi_field_crops"
                    and ai_extraction_payload["roi_template_type"] == "blue_structured_card",
                    "AI extraction response should disclose ROI-based extraction mode and template.",
                )
                with db.connection() as conn:
                    ai_record = conn.execute(
                        "SELECT raw_payload FROM parse_result WHERE parse_id = ?",
                        (ai_extraction_payload["parse_id"],),
                    ).fetchone()
                ai_raw_payload = json.loads(ai_record["raw_payload"])
                assert_true(
                    ai_raw_payload["extractionImageMode"] == "roi_field_crops"
                    and ai_raw_payload["uncertainFields"] == ["dob_normalized"]
                    and ai_raw_payload["symbolConfusions"] == ["F vs \u2640 resolved from Sex ROI"]
                    and ai_raw_payload["rawVisibleTextLines"],
                    "AI extraction should persist structured uncertainty, symbol confusion, raw visible text, and ROI mode evidence.",
                )
                ai_reviews = [
                    item for item in client.get("/api/review-items").json()
                    if item["parse_id"] == ai_extraction_payload["parse_id"]
                ]
                assert_true(
                    any(item["issue"] == "AI-extracted photo transcription needs review" for item in ai_reviews),
                    "AI extraction should create a review item distinct from manual transcription.",
                )
                ai_workbench = client.get("/api/photo-review-workbench").json()
                ai_workbench_row = next(item for item in ai_workbench["rows"] if item["photo_id"] == ai_photo_payload["photo_id"])
                assert_true(
                    ai_workbench_row["manual_parse_id"] == ai_extraction_payload["parse_id"]
                    and ai_workbench_row["note_line_count"] == 2,
                    "Photo workbench should treat saved AI extraction as reviewable transcription progress.",
                )
                photos = client.get("/api/photos").json()
                assert_true(
                    photos[0]["open_review_count"] >= 1 and photos[0]["latest_parse_id"],
                    "Photo list should expose review candidate context.",
                )
                photo_reviews = [
                    item for item in client.get("/api/review-items").json()
                    if item["photo_id"] == photo_payload["photo_id"]
                ]
                assert_true(photo_reviews, "Photo review candidate should be visible in Review Queue.")
                assert_true(
                    photo_reviews[0]["assigned_role"] == "Colony Reviewer"
                    and photo_reviews[0]["priority"] in {"medium", "high", "low"},
                    "Review Queue API should expose persona assignment and priority for operational queues.",
                )
                assert_true(
                    "manual cage-card transcription" in photo_reviews[0]["issue"],
                    "Photo review candidate should require manual transcription.",
                )
                manual_transcription = client.post(
                    f"/api/photos/{photo_payload['photo_id']}/manual-transcription",
                    json={
                        "card_type": "Separated",
                        "raw_strain": "ApoM Tg/Tg",
                        "sex_raw": "♀",
                        "id_raw": "MT",
                        "dob_raw": "25.10.20-28",
                        "mouse_count": "2 total",
                        "notes": [
                            {"raw": "MT321 R'", "strike": "none"},
                            {"raw": "MT322 L'", "strike": "none"},
                            {"raw": "1 2 3", "strike": "none", "meaning": "unlabeled_numeric_note"},
                        ],
                    },
                )
                assert_true(manual_transcription.status_code == 200, f"Could not create manual transcription: {manual_transcription.text}")
                transcription_payload = manual_transcription.json()
                assert_true(
                    transcription_payload["boundary"] == "parsed or intermediate result",
                    "Manual photo transcription should stay parsed/intermediate.",
                )
                assert_true(
                    transcription_payload["created_note_items"] == 3
                    and transcription_payload["created_mouse_candidates"] == 0,
                    "Manual photo transcription should create note evidence without canonical mouse candidates, including numeric-only labels.",
                )
                manual_review_items = [
                    item for item in client.get("/api/review-items").json()
                    if item["parse_id"] == transcription_payload["parse_id"]
                ]
                assert_true(
                    any(
                        item["issue"] == "Unlabeled numeric note needs review"
                        and item["current_value"] == "1, 2, 3"
                        and "numeric-only note lines" in item["review_reason"]
                        for item in manual_review_items
                    ),
                    "Numeric-only manual note lines should open one grouped review item instead of becoming mouse IDs.",
                )
                assert_true(
                    transcription_payload["resolved_photo_review_items"] >= 1,
                    "Manual transcription should resolve the initial photo-level review candidate.",
                )
                transcribed_workbench = client.get("/api/photo-review-workbench").json()
                transcribed_row = next(item for item in transcribed_workbench["rows"] if item["photo_id"] == photo_payload["photo_id"])
                assert_true(
                    transcribed_row["manual_parse_id"] == transcription_payload["parse_id"]
                    and transcribed_row["note_line_count"] == 3
                    and transcribed_row["next_action"] == "resolve_photo_reviews",
                    "Photo review workbench should expose manual transcription progress and remaining review work.",
                )
                numeric_review = next(
                    item for item in manual_review_items
                    if item["issue"] == "Unlabeled numeric note needs review"
                )
                assert_true(
                    numeric_review["note_item_id"] and numeric_review["review_note_raw_line"] == "1 2 3",
                    "Numeric note review items should expose the exact note-line anchor.",
                )
                assert_true(
                    numeric_review["card_snapshot_id"] == transcription_payload["card_snapshot_id"]
                    and numeric_review["review_card_type"].lower() == "separated"
                    and numeric_review["review_note_summary"]["note_count"] == 3
                    and numeric_review["image_url"].endswith("/image"),
                    "Numeric note review detail should expose linked card snapshot and raw photo evidence.",
                )
                assert_true(
                    numeric_review["image_url"].endswith("/image")
                    and numeric_review["card_snapshot_id"] == transcription_payload["card_snapshot_id"]
                    and numeric_review["review_note_summary"]["needs_review_count"] >= 1,
                    "Review Queue items should expose raw photo preview URLs and card snapshot evidence context.",
                )
                numeric_resolution = client.post(
                    f"/api/review-items/{numeric_review['review_id']}/resolve",
                    json={
                        "resolution_note": "Confirmed numeric-only line is a reviewed count note, not a mouse ID.",
                        "resolved_value": "3 temporary labels",
                        "note_item_id": numeric_review["note_item_id"],
                        "note_label_decision": "count_note",
                        "note_label_count": 3,
                    },
                )
                assert_true(numeric_resolution.status_code == 200, f"Could not resolve numeric note label: {numeric_resolution.text}")
                numeric_resolution_payload = numeric_resolution.json()
                assert_true(
                    numeric_resolution_payload["note_label_update"]["boundary"] == "parsed or intermediate result"
                    and numeric_resolution_payload["note_label_update"]["decision"] == "count_note",
                    "Resolving a numeric note label should stay in the parsed/intermediate layer.",
                )
                card_snapshots = client.get("/api/card-snapshots").json()
                assert_true(
                    card_snapshots
                    and card_snapshots[0]["boundary"] == "parsed or intermediate result"
                    and any(item["card_snapshot_id"] == transcription_payload["card_snapshot_id"] for item in card_snapshots),
                    "Manual transcription should create a parsed card snapshot that remains non-canonical evidence.",
                )
                reviewed_snapshot = next(
                    item for item in card_snapshots
                    if item["card_snapshot_id"] == transcription_payload["card_snapshot_id"]
                )
                assert_true(
                    reviewed_snapshot["note_summary"].get("count_note_total") == 3
                    and reviewed_snapshot["status"] in {"review", "reviewed"},
                    "Resolving a numeric note label should refresh the card snapshot note summary.",
                )
                corrected_note_items = client.get("/api/note-items").json()
                corrected_numeric_note = next(
                    item for item in corrected_note_items
                    if item["note_item_id"] == numeric_review["note_item_id"]
                )
                assert_true(
                    corrected_numeric_note["parsed_type"] == "count_note"
                    and corrected_numeric_note["parsed_count"] == 3
                    and corrected_numeric_note["needs_review"] == 0,
                    "Numeric note label correction should update the note item without inventing a mouse record.",
                )
                numeric_corrections = [
                    item for item in client.get("/api/corrections").json()
                    if item["entity_type"] == "note_item"
                    and item["entity_id"] == numeric_review["note_item_id"]
                    and item["field_name"] == "parsed_label"
                ]
                assert_true(
                    numeric_corrections,
                    "Numeric note label correction should preserve before/after values in correction_log.",
                )
                idempotent_photo_reviews = client.post("/api/photos/review-candidates")
                assert_true(idempotent_photo_reviews.status_code == 200, "Could not run photo review candidate backfill.")
                assert_true(
                    idempotent_photo_reviews.json()["existing_review_candidates"] >= 1,
                    "Photo review candidate backfill should be idempotent.",
                )
                reconciliation = client.get("/api/evidence-reconciliation").json()
                assert_true(reconciliation["boundary"] == "export or view", "Evidence reconciliation should remain a view.")
                assert_true(reconciliation["total_photos"] >= 1, "Evidence reconciliation should count uploaded photos.")
                assert_true(reconciliation["legacy_rows"] >= 2, "Evidence reconciliation should count legacy workbook rows.")
                assert_true(reconciliation["manual_transcriptions"] >= 1, "Evidence reconciliation should count manual photo transcriptions.")
                assert_true(
                    reconciliation["source_priority"][0] == "raw source photo",
                    "Evidence reconciliation should prioritize raw photos over predecessor Excel views.",
                )
                mice_before_comparison = client.get("/api/mice").json()
                comparison = client.get("/api/evidence-comparison")
                assert_true(comparison.status_code == 200, f"Could not load evidence comparison: {comparison.text}")
                comparison_payload = comparison.json()
                assert_true(comparison_payload["boundary"] == "export or view", "Evidence comparison should remain a view.")
                assert_true(
                    comparison_payload["manual_transcription_count"] >= 1,
                    "Evidence comparison should count manual photo transcriptions.",
                )
                assert_true(
                    comparison_payload["legacy_candidate_count"] >= 1,
                    "Evidence comparison should count predecessor Excel candidates.",
                )
                assert_true(
                    comparison_payload["comparison_count"] >= 1,
                    "Evidence comparison should produce manual-vs-legacy comparison rows.",
                )
                first_comparison = comparison_payload["comparisons"][0]
                assert_true(
                    {"source_layer", "manual_summary", "status", "detail", "review_required", "review_status"}.issubset(first_comparison),
                    "Evidence comparison rows should expose review-view context.",
                )
                assert_true(
                    first_comparison["manual_summary"].get("sex") == "♀"
                    and first_comparison["manual_summary"].get("id") == "MT",
                    "Evidence comparison should preserve raw Sex and I.D fields from manual transcription.",
                )
                assert_true(
                    first_comparison["review_status"] in {"not_created", "not_required", "open", "resolved"},
                    "Evidence comparison rows should expose actionable review state.",
                )
                assert_true(
                    client.get("/api/mice").json() == mice_before_comparison,
                    "Evidence comparison should not write canonical mouse state.",
                )
                scoped_comparison = next(
                    item for item in comparison_payload["comparisons"]
                    if item["review_required"]
                )
                comparison_review = client.post(
                    "/api/evidence-comparison/review-candidates",
                    json={"manual_parse_id": scoped_comparison["manual_parse_id"]},
                )
                assert_true(
                    comparison_review.status_code == 200,
                    f"Could not create evidence comparison review candidates: {comparison_review.text}",
                )
                comparison_review_payload = comparison_review.json()
                assert_true(
                    comparison_review_payload["boundary"] == "review item",
                    "Evidence comparison review candidate creation should stay in the review layer.",
                )
                assert_true(
                    comparison_review_payload["scope"]["manual_parse_id"] == scoped_comparison["manual_parse_id"]
                    and comparison_review_payload["scope"]["matched_comparisons"] == 1
                    and comparison_review_payload["created_review_items"] == 1,
                    "Scoped evidence comparison review creation should only create the targeted manual parse review.",
                )
                comparison_review_ids = set(comparison_review_payload["review_ids"])
                review_items_after_comparison = client.get("/api/review-items").json()
                comparison_review_items = [
                    item
                    for item in review_items_after_comparison
                    if item["review_id"] in comparison_review_ids
                ]
                assert_true(
                    comparison_review_items,
                    (
                        "Evidence comparison review candidates should be visible in Review Queue. "
                        f"Expected ids: {sorted(comparison_review_ids)}; "
                        f"recent ids: {[item['review_id'] for item in review_items_after_comparison[:8]]}"
                    ),
                )
                comparison_after_review = client.get("/api/evidence-comparison").json()
                comparison_review_states = {
                    item["review_id"]: item["review_status"]
                    for item in comparison_after_review["comparisons"]
                    if item["review_id"] in comparison_review_ids
                }
                assert_true(
                    comparison_review_states and set(comparison_review_states.values()) == {"open"},
                    "Evidence comparison rows should reflect Review Queue state after candidate creation.",
                )
                idempotent_comparison_review = client.post(
                    "/api/evidence-comparison/review-candidates",
                    json={"manual_parse_id": scoped_comparison["manual_parse_id"]},
                )
                assert_true(
                    idempotent_comparison_review.json()["created_review_items"] == 0,
                    "Scoped evidence comparison review candidate creation should be idempotent.",
                )
                batch_scoped_comparison_review = client.post(
                    "/api/evidence-comparison/review-candidates",
                    json={"upload_batch_id": upload_batch_payload["upload_batch_id"]},
                )
                assert_true(
                    batch_scoped_comparison_review.json()["scope"]["upload_batch_id"] == upload_batch_payload["upload_batch_id"]
                    and batch_scoped_comparison_review.json()["scope"]["matched_comparisons"] == 1
                    and batch_scoped_comparison_review.json()["existing_review_items"] == 1,
                    "Upload-batch scoped comparison review creation should only touch comparisons from the selected batch.",
                )
                assert_true(
                    client.get("/api/mice").json() == mice_before_comparison,
                    "Evidence comparison review candidate creation should not write canonical mouse state.",
                )
                draft_candidates_before = client.get("/api/canonical-candidates").json()
                mapped_review_id = comparison_review_payload["review_ids"][0]
                mapped_review_item = next(item for item in comparison_review_items if item["review_id"] == mapped_review_id)
                mapped_review = client.post(
                    f"/api/review-items/{mapped_review_id}/resolve",
                    json={
                        "resolution_note": "Map reviewed photo-vs-Excel comparison into a draft candidate.",
                        "resolved_value": "draft candidate",
                        "legacy_decision": "map_to_canonical_candidate",
                    },
                )
                assert_true(mapped_review.status_code == 200, f"Could not map comparison review: {mapped_review.text}")
                mapped_payload = mapped_review.json()
                assert_true(
                    mapped_payload["canonical_candidate_id"],
                    "Mapping a comparison review should create a canonical candidate draft.",
                )
                draft_candidates_after = client.get("/api/canonical-candidates").json()
                assert_true(
                    len(draft_candidates_after) == len(draft_candidates_before) + 1,
                    "Canonical candidate draft list should include the mapped comparison review.",
                )
                created_candidate = next(
                    item for item in draft_candidates_after
                    if item["candidate_id"] == mapped_payload["canonical_candidate_id"]
                )
                assert_true(
                    created_candidate["status"] == "draft"
                    and created_candidate["review_id"] == mapped_review_id,
                    "Canonical candidate should remain a draft linked to the resolved review.",
                )
                assert_true(
                    created_candidate["boundary"] == "review item"
                    and created_candidate["source_layer"] == "review item",
                    "Canonical candidate drafts should be explicitly classified as review-layer records.",
                )
                assert_true(
                    created_candidate["parse_id"] == mapped_review_item["parse_id"],
                    "Canonical candidate drafts should retain parse traceability.",
                )
                assert_true(
                    client.get("/api/mice").json() == mice_before_comparison,
                    "Mapping a comparison review should not write canonical mouse state.",
                )
                preview_candidate = client.get(f"/api/canonical-candidates/{mapped_payload['canonical_candidate_id']}/apply-preview")
                assert_true(preview_candidate.status_code == 200, f"Could not preview canonical candidate apply: {preview_candidate.text}")
                preview_payload = preview_candidate.json()
                assert_true(
                    preview_payload["boundary"] == "export or view",
                    "Canonical candidate apply preview should remain a view.",
                )
                assert_true(
                    preview_payload["summary"]["new_mouse_rows"] == 2
                    and preview_payload["summary"]["events"] == 2
                    and preview_payload["blocked"] is False,
                    "Canonical candidate apply preview should show mouse/event writes before applying.",
                )
                assert_true(
                    client.get("/api/mice").json() == mice_before_comparison,
                    "Canonical candidate apply preview should not write mouse state.",
                )
                apply_candidate = client.post(f"/api/canonical-candidates/{mapped_payload['canonical_candidate_id']}/apply")
                assert_true(apply_candidate.status_code == 200, f"Could not apply canonical candidate draft: {apply_candidate.text}")
                apply_payload = apply_candidate.json()
                assert_true(
                    apply_payload["boundary"] == "canonical structured state",
                    "Applying a draft should be explicitly classified as canonical structured state.",
                )
                assert_true(
                    apply_payload["created_mice"] == 2 and apply_payload["created_events"] == 2,
                    "Applying the comparison draft should create mouse records and events from parsed note lines.",
                )
                applied_candidates = client.get("/api/canonical-candidates").json()
                applied_candidate = next(
                    item for item in applied_candidates
                    if item["candidate_id"] == mapped_payload["canonical_candidate_id"]
                )
                assert_true(applied_candidate["status"] == "applied", "Applied canonical candidate should expose applied status.")
                applied_mice = client.get("/api/mice").json()
                applied_mt321 = next((mouse for mouse in applied_mice if mouse["display_id"] == "MT321"), None)
                applied_mt322 = next((mouse for mouse in applied_mice if mouse["display_id"] == "MT322"), None)
                assert_true(
                    applied_mt321 is not None and applied_mt322 is not None,
                    "Applied canonical candidate should create reviewed mouse records from note lines.",
                )
                assert_true(
                    applied_mt321["source_note_item_id"] and applied_mt322["source_note_item_id"],
                    "Applied canonical mouse records should retain source note-line anchors.",
                )
                assert_true(
                    applied_mt321["last_verified_at"] == apply_payload["applied_at"]
                    and applied_mt322["last_verified_at"] == apply_payload["applied_at"],
                    "Canonical apply should stamp mouse current state with last_verified_at.",
                )
                applied_events = [
                    item
                    for item in client.get("/api/mouse-events").json()
                    if item["related_entity_id"] == mapped_payload["canonical_candidate_id"]
                    and item["event_type"] == "canonical_candidate_applied"
                ]
                assert_true(
                    len(applied_events) == 2,
                    "Applying a canonical candidate should create source-backed mouse events.",
                )
                idempotent_apply = client.post(f"/api/canonical-candidates/{mapped_payload['canonical_candidate_id']}/apply")
                assert_true(
                    idempotent_apply.status_code == 409,
                    "Re-applying an applied canonical candidate should be blocked instead of silently rewriting state.",
                )
                candidate_audit = client.get(f"/api/canonical-candidates/{mapped_payload['canonical_candidate_id']}/audit")
                assert_true(candidate_audit.status_code == 200, f"Could not audit applied canonical candidate: {candidate_audit.text}")
                candidate_audit_payload = candidate_audit.json()
                assert_true(
                    candidate_audit_payload["boundary"] == "export or view",
                    "Canonical candidate audit should remain a view.",
                )
                assert_true(
                    candidate_audit_payload["can_void"] is True
                    and candidate_audit_payload["summary"]["applied_mouse_count"] == 2
                    and candidate_audit_payload["summary"]["event_count"] >= 2,
                    "Applied candidate audit should expose linked mouse records and events before void.",
                )
                applied_action = next(
                    action for action in candidate_audit_payload["actions"]
                    if action["action_type"] == "canonical_candidate_applied"
                )
                assert_true(
                    applied_action["after_value"]["review_id"] == mapped_review_id
                    and applied_action["after_value"]["source_note_item_ids"]
                    and applied_action["after_value"]["raw_note_lines"],
                    "Canonical apply action should preserve review, source note, and raw line evidence.",
                )
                created = client.post(
                    "/api/assigned-strains",
                    json={
                        "display_name": "ApoM Tg/Tg",
                        "aliases": ["ApoMtg/tg", "ApoM"],
                        "source_type": "manual",
                    },
                )
                assert_true(created.status_code == 200, "Could not create assigned strain scope.")
                payload = created.json()
                assert_true(payload["active"] is True, "Created assigned strain should be active.")
                assert_true("ApoMtg/tg" in payload["aliases"], "Assigned strain aliases were not preserved.")

                rows = client.get("/api/assigned-strains").json()
                assert_true(len(rows) == 1, "Assigned strain list did not return the created scope.")
                assert_true(rows[0]["display_name"] == "ApoM Tg/Tg", "Assigned strain display name changed.")

                deactivated = client.post(f"/api/assigned-strains/{payload['assigned_strain_id']}/deactivate")
                assert_true(deactivated.status_code == 200, "Could not deactivate assigned strain scope.")
                rows = client.get("/api/assigned-strains").json()
                assert_true(rows[0]["active"] is False, "Deactivated assigned strain stayed active.")

                mice_before_distribution = client.get("/api/mice").json()
                distribution_fixture = json.loads(
                    (ROOT / "fixtures" / "sample_distribution_import.json").read_text(encoding="utf-8")
                )
                distribution_import = client.post("/api/distribution-imports", json=distribution_fixture)
                assert_true(distribution_import.status_code == 200, "Could not import distribution assignment JSON.")
                distribution_payload = distribution_import.json()
                assert_true(
                    distribution_payload["boundary"] == "parsed or intermediate result",
                    "Distribution import should stay non-canonical.",
                )
                assert_true(distribution_payload["stored_rows"] >= 3, "Distribution import stored too few rows.")
                imports = client.get("/api/distribution-imports").json()
                imported_fixture = next(
                    (
                        item
                        for item in imports
                        if item["distribution_import_id"] == distribution_payload["distribution_import_id"]
                    ),
                    None,
                )
                assert_true(imported_fixture is not None, f"Distribution import list did not return the new import: {imports!r}")
                assert_true(
                    any(row["mating_type_raw"] == "GFAP Cre; S1PR1 fl/fl" for row in imported_fixture["rows"]),
                    "Distribution assignment rows did not preserve candidate mating type.",
                )
                assert_true(
                    client.get("/api/mice").json() == mice_before_distribution,
                    "Distribution import must not create canonical mouse records.",
                )

                with db.connection() as conn:
                    action_count = conn.execute(
                        "SELECT COUNT(*) AS count FROM action_log WHERE target_id = ?",
                        (payload["assigned_strain_id"],),
                    ).fetchone()["count"]
                assert_true(action_count == 2, "Assigned strain changes should be logged.")

                outside_import = client.post("/api/fixtures/import-sample")
                assert_true(outside_import.status_code == 200, "Could not import fixture without active assigned scope.")
                outside_reviews = client.get("/api/review-items").json()
                assert_true(
                    any(item["issue"] == "Outside assigned strain scope" for item in outside_reviews),
                    "Fixture import without active scope should create outside-scope review items.",
                )

                client.post(
                    "/api/assigned-strains",
                    json={
                        "display_name": "ApoM Tg/Tg",
                        "aliases": ["ApoMtg/tg", "ApoM"],
                        "source_type": "manual",
                    },
                )
                imported = client.post("/api/fixtures/import-sample")
                assert_true(imported.status_code == 200, "Could not import sample fixture through local API.")
                imported_payload = imported.json()
                assert_true(
                    imported_payload["created_or_updated_note_items"] >= 10,
                    "Fixture import should persist parsed note item evidence.",
                )
                assert_true(
                    imported_payload["created_or_updated_mouse_candidates"] >= 3,
                    "Fixture import should create safe mouse candidates from accepted separated rows.",
                )
                mouse_events = client.get("/api/mouse-events").json()
                assert_true(len(mouse_events) >= 3, "Mouse candidates should create mouse event history.")
                assert_true(
                    any(event["event_type"] == "note_added" for event in mouse_events),
                    "Mouse event history should include source-backed note events.",
                )
                review_items = client.get("/api/review-items").json()
                assert_true(len(review_items) >= 4, "Fixture import should create review and validation items.")
                assert_true(
                    any(item["evidence_preview"] and item["note_line_count"] >= 1 for item in review_items),
                    "Review items should expose source note evidence context.",
                )
                assert_true(
                    any(item["source_name"] or item["photo_id"] or item["original_filename"] for item in review_items),
                    "Review items should expose source record or photo context.",
                )
                partial_correction = client.post(
                    f"/api/review-items/{review_items[0]['review_id']}/resolve",
                    json={
                        "resolution_note": "Incomplete correction payload should not be accepted.",
                        "correction_entity_type": "review_item",
                    },
                )
                assert_true(partial_correction.status_code == 400, "Partial review correction payload should be rejected.")
                assert_true(
                    not any(item["issue"] == "Outside assigned strain scope" for item in review_items),
                    "Assigned ApoM scope should clear stale outside-scope review items for now-accepted rows.",
                )
                assert_true(
                    any(item["parse_id"] == "FIXTURE-COUNT-MISMATCH" and item["issue"] == "Count mismatch" for item in review_items),
                    "Count mismatch fixture should create a backend review item.",
                )
                assert_true(
                    any(item["parse_id"] == "FIXTURE-DUPLICATE-ACTIVE" and item["issue"] == "Duplicate active mouse" for item in review_items),
                    "Duplicate active fixture should create a backend review item.",
                )
                assert_true(
                    any(item["parse_id"] == "FIXTURE-EAR-LABEL-CHECK" and item["issue"] == "Ear label needs review" for item in review_items),
                    "Ambiguous ear label fixture should create a note-level review item.",
                )
                count_review = next(item for item in review_items if item["parse_id"] == "FIXTURE-COUNT-MISMATCH")
                resolved = client.post(
                    f"/api/review-items/{count_review['review_id']}/resolve",
                    json={
                        "resolution_note": "Reviewed count mismatch in source note lines.",
                        "resolved_value": count_review["suggested_value"],
                        "correction_entity_type": "review_item",
                        "correction_entity_id": count_review["review_id"],
                        "correction_field_name": "mouse_count",
                        "correction_before_value": count_review["current_value"],
                        "correction_after_value": count_review["suggested_value"],
                    },
                )
                assert_true(resolved.status_code == 200, "Could not resolve review item.")
                resolved_payload = resolved.json()
                assert_true(resolved_payload["correction_id"], "Review resolution with correction values should create a correction log entry.")
                assert_true(resolved_payload["audit_url"].endswith("/audit"), "Review resolution should return an audit endpoint.")
                resolved_items = client.get("/api/review-items").json()
                resolved_count_review = next(item for item in resolved_items if item["review_id"] == count_review["review_id"])
                assert_true(resolved_count_review["status"] == "resolved", "Resolved review item stayed open.")
                review_audit = client.get(resolved_payload["audit_url"])
                assert_true(review_audit.status_code == 200, "Review audit endpoint should load after resolution.")
                review_audit_payload = review_audit.json()
                review_resolved_action = next(
                    action for action in review_audit_payload["actions"]
                    if action["action_type"] == "review_resolved"
                )
                assert_true(
                    review_audit_payload["source_layer"] == "export or view"
                    and review_audit_payload["summary"]["action_count"] >= 1
                    and review_audit_payload["summary"]["correction_count"] == 1,
                    "Review audit should expose action and correction trace as a view.",
                )
                assert_true(
                    review_resolved_action["after_value"]["correction_id"] == resolved_payload["correction_id"]
                    and review_resolved_action["after_value"]["source_parse_id"] == count_review["parse_id"]
                    and "boundary" in review_resolved_action["after_value"],
                    "Review resolved action should preserve correction and source context.",
                )
                with db.connection() as conn:
                    review_action_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM action_log
                        WHERE action_type = 'review_resolved' AND target_id = ?
                        """,
                        (count_review["review_id"],),
                    ).fetchone()["count"]
                    correction_row = conn.execute(
                        """
                        SELECT correction_id, entity_type, entity_id, field_name,
                               before_value, after_value, reason, review_id
                        FROM correction_log
                        WHERE correction_id = ?
                        """,
                        (resolved_payload["correction_id"],),
                    ).fetchone()
                    correction_action_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM action_log
                        WHERE action_type = 'correction_applied'
                          AND target_id = ?
                        """,
                        (count_review["review_id"],),
                    ).fetchone()["count"]
                assert_true(review_action_count == 1, "Review resolution should create an action log entry.")
                assert_true(
                    correction_row is not None
                    and correction_row["before_value"] == count_review["current_value"]
                    and correction_row["after_value"] == count_review["suggested_value"]
                    and correction_row["review_id"] == count_review["review_id"],
                    "Review resolution correction should preserve before/after values and review linkage.",
                )
                assert_true(correction_action_count == 1, "Review resolution correction should create a correction action log entry.")
                duplicate_resolve = client.post(
                    f"/api/review-items/{count_review['review_id']}/resolve",
                    json={"resolution_note": "Duplicate review resolution should be blocked."},
                )
                assert_true(duplicate_resolve.status_code == 409, "Resolved review items should not be resolved again.")
                duplicate_active_review = next(item for item in resolved_items if item["parse_id"] == "FIXTURE-DUPLICATE-ACTIVE")
                partial_correction = client.post(
                    f"/api/review-items/{duplicate_active_review['review_id']}/resolve",
                    json={
                        "resolution_note": "Incomplete correction metadata should be rejected.",
                        "correction_entity_type": "review_item",
                    },
                )
                assert_true(partial_correction.status_code == 400, "Partial correction metadata should not resolve a review item.")
                still_open_reviews = client.get("/api/review-items").json()
                assert_true(
                    any(
                        item["review_id"] == duplicate_active_review["review_id"]
                        and item["status"] == "open"
                        for item in still_open_reviews
                    ),
                    "Rejected partial correction should leave the review item open.",
                )
                note_items = client.get("/api/note-items").json()
                mice = client.get("/api/mice").json()
                assert_true(
                    any(item["raw_line_text"] == "MT321 R'" and item["parsed_ear_label_code"] == "R_PRIME" for item in note_items),
                    "Note item API should expose parsed ear label evidence.",
                )
                assert_true(
                    any(
                        item["raw_line_text"] == "1 2 3"
                        and item["parsed_type"] == "count_note"
                        and item["parsed_count"] == 3
                        and item["needs_review"] == 0
                        for item in note_items
                    ),
                    "Reviewed numeric-only post-separation note lines should keep raw evidence without becoming mouse IDs.",
                )
                assert_true(
                    any(
                        item["raw_line_text"] == "MT399 R0"
                        and item["parsed_ear_label_code"] == "R_CIRCLE"
                        and item["parsed_ear_label_review_status"] == "check"
                        and item["needs_review"] == 1
                        for item in note_items
                    ),
                    "Ambiguous ear label note should stay reviewable with raw evidence.",
                )
                assert_true(
                    any(
                        mouse["display_id"] == "MT399"
                        and mouse["ear_label_raw"] == "R0"
                        and mouse["ear_label_code"] is None
                        and mouse["ear_label_review_status"] == "check"
                        for mouse in mice
                    ),
                    "Mouse candidate should not accept an uncertain normalized ear label.",
                )
                ambiguous_ear_note = next(item for item in note_items if item["raw_line_text"] == "MT399 R0")
                ear_resolution = client.post(
                    f"/api/review-items/review_ear_{ambiguous_ear_note['note_item_id']}/resolve",
                    json={
                        "resolution_note": "Checked the source card; circle-vs-zero mark is R circle.",
                        "resolved_value": "R_CIRCLE",
                        "note_item_id": ambiguous_ear_note["note_item_id"],
                        "ear_label_code": "R_CIRCLE",
                    },
                )
                assert_true(ear_resolution.status_code == 200, "Ear label review should resolve from a bounded select value.")
                ear_resolution_payload = ear_resolution.json()
                assert_true(
                    ear_resolution_payload["ear_label_update"]["ear_label_code"] == "R_CIRCLE"
                    and ear_resolution_payload["ear_label_update"]["boundary"] == "parsed or intermediate result",
                    "Ear label resolution should update parsed evidence without becoming canonical state.",
                )
                updated_note_items = client.get("/api/note-items").json()
                assert_true(
                    any(
                        item["note_item_id"] == ambiguous_ear_note["note_item_id"]
                        and item["raw_line_text"] == "MT399 R0"
                        and item["parsed_ear_label_code"] == "R_CIRCLE"
                        and item["parsed_ear_label_review_status"] == "user_corrected"
                        and item["needs_review"] == 0
                        for item in updated_note_items
                    ),
                    "Ear label correction should preserve raw note text while clearing the note-level review flag.",
                )
                assert_true(
                    any(mouse["display_id"] == "MT323" and mouse["status"] == "moved" for mouse in mice),
                    "Mouse API should expose moved candidate from single-struck note line.",
                )
                cage = client.post(
                    "/api/cages",
                    json={"cage_label": "C-014", "location": "Room A / Rack 2", "cage_type": "holding"},
                )
                assert_true(cage.status_code == 200, "Could not create cage registry entry.")
                cage_payload = cage.json()
                assert_true(cage_payload["source_record_id"], "Cage creation should create source evidence.")
                moved_mouse = next(
                    mouse for mouse in mice
                    if mouse["display_id"] == "MT321" and mouse["status"] == "active"
                )
                moved_to_cage = client.post(
                    f"/api/mice/{moved_mouse['mouse_id']}/move-cage",
                    json={"cage_id": cage_payload["cage_id"], "note": "Verified cage movement flow."},
                )
                assert_true(moved_to_cage.status_code == 200, "Could not move mouse to cage.")
                moved_payload = moved_to_cage.json()
                assert_true(moved_payload["event_id"], "Cage move should create a mouse event.")
                cage_rows = client.get("/api/cages").json()
                assert_true(
                    any(row["cage_label"] == "C-014" and row["active_mouse_count"] == 1 for row in cage_rows),
                    "Cage list should show active mouse count after assignment.",
                )
                mice_after_cage = client.get("/api/mice", params={"query": "C-014"}).json()
                assert_true(
                    any(mouse["display_id"] == "MT321" and mouse["current_cage_label"] == "C-014" for mouse in mice_after_cage),
                    "Mouse search should find current cage assignment.",
                )
                cage_events = client.get("/api/mouse-events").json()
                assert_true(
                    any(event["event_type"] == "moved" and event["related_entity_id"] == cage_payload["cage_id"] for event in cage_events),
                    "Cage movement should be present in mouse event history.",
                )
                audit_trace = client.get(f"/api/mice/{moved_mouse['mouse_id']}/audit-trace")
                assert_true(audit_trace.status_code == 200, "Could not load mouse audit trace.")
                audit_payload = audit_trace.json()
                assert_true(audit_payload["source_layer"] == "export or view", "Audit trace should stay in the export/view layer.")
                assert_true(audit_payload["mouse"]["display_id"] == "MT321", "Audit trace should return the requested mouse.")
                audit_categories = {item["category"] for item in audit_payload["timeline"]}
                assert_true({"note_line", "mouse_event"}.issubset(audit_categories), "Audit trace should include source note lines and mouse events.")
                assert_true(
                    any(event["event_type"] == "moved" for event in audit_payload["events"]),
                    "Audit trace should include cage movement events.",
                )
                filtered_mice = client.get("/api/mice", params={"query": "MT321"}).json()
                assert_true(
                    filtered_mice and all("MT321" in mouse["display_id"] for mouse in filtered_mice),
                    "Mouse search filter should narrow candidate records by display ID.",
                )
                search_payload = client.get("/api/search", params={"query": "PV-Cre"}).json()
                assert_true(search_payload["query"] == "PV-Cre", "Search endpoint should echo the active query.")
                assert_true(
                    any(strain["strain_name"] == "PV-Cre" for strain in search_payload["strains"]),
                    "Search endpoint should include strain registry matches.",
                )
                csv_response = client.get("/api/exports/mice.csv", params={"query": "MT321"})
                assert_true(csv_response.status_code == 200, "Mouse CSV export endpoint failed.")
                assert_true(
                    csv_response.headers["content-type"].startswith("text/csv"),
                    "Mouse CSV export should return CSV content.",
                )
                csv_text = csv_response.text
                assert_true("display_id" in csv_text and "MT321" in csv_text, "Mouse CSV export is missing expected rows.")
                assert_true("MT323" not in csv_text, "Filtered mouse CSV export should exclude non-matching mice.")
                export_log = client.get("/api/export-log").json()
                assert_true(export_log, "CSV generation should create an export log entry.")
                assert_true(export_log[0]["export_type"] == "mouse_csv", "Export log should record CSV export type.")
                assert_true(export_log[0]["filename"] == "mouse_records_filtered.csv", "Export log should preserve generated filename.")
                assert_true(export_log[0]["query"] == "MT321", "Export log should preserve the export filter.")
                assert_true(export_log[0]["row_count"] == len(filtered_mice), "Export log should preserve exported row count.")
                assert_true(export_log[0]["source_layer"] == "export or view", "Export log should stay in the export/view layer.")
                assert_true("last_verified_at" in csv_text, "Mouse CSV export should include last verified timestamps.")
                assert_true(
                    export_log[0]["generated_role"] == "Data / Export Manager"
                    and export_log[0]["generated_by"] == "local_user",
                    "Export log should expose persona ownership for generated export views.",
                )
                with db.connection() as conn:
                    conn.execute(
                        """
                        INSERT INTO review_queue
                            (review_id, parse_id, severity, issue, current_value, suggested_value,
                             review_reason, status, created_at)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                        """,
                        (
                            "review_export_blocker_after_scoped_comparison",
                            photo_payload["review_candidate"]["parse_id"],
                            "High",
                            "Export readiness blocker after scoped comparison review",
                            "scoped comparison verified",
                            "resolve before final export",
                            "Keeps final export readiness checks covered after the scoped comparison review is resolved.",
                            "open",
                        ),
                    )
                blocked_export = client.get("/api/exports/mice.csv", params={"query": "MT321", "require_ready": "true"})
                assert_true(blocked_export.status_code == 409, "Final CSV export should be blocked by open review items.")
                blocked_separation_xlsx = client.get("/api/exports/separation.xlsx")
                assert_true(blocked_separation_xlsx.status_code == 409, "Final separation XLSX export should be blocked by open review items.")
                blocked_animal_xlsx = client.get("/api/exports/animal-sheet.xlsx")
                assert_true(blocked_animal_xlsx.status_code == 409, "Final animal sheet XLSX export should be blocked by open review items.")
                blocked_payload = blocked_export.json()["detail"]
                assert_true(
                    blocked_payload["source_layer"] == "export or view" and blocked_payload["blocked_review_count"] > 0,
                    "Blocked final export should report export/view layer and blocker count.",
                )
                assert_true(
                    blocked_payload["review_blockers"]
                    and {
                        "review_id",
                        "issue",
                        "severity",
                        "review_reason",
                        "evidence_preview",
                        "note_line_count",
                        "review_check_targets",
                    }.issubset(blocked_payload["review_blockers"][0]),
                    "Blocked final CSV export should include actionable review blocker details.",
                )
                assert_true(
                    blocked_payload["review_blockers"][0]["review_check_targets"],
                    "Blocked final CSV export should include focused check targets for review blockers.",
                )
                blocked_separation_payload = blocked_separation_xlsx.json()["detail"]
                blocked_animal_payload = blocked_animal_xlsx.json()["detail"]
                assert_true(
                    blocked_separation_payload["review_blockers"]
                    and blocked_animal_payload["review_blockers"],
                    "Blocked workbook exports should include review blocker previews.",
                )
                blocked_logs = client.get("/api/export-log").json()
                blocked_log = next(
                    (item for item in blocked_logs if item["export_type"] == "mouse_csv" and item["status"] == "blocked"),
                    None,
                )
                assert_true(blocked_log is not None, "Blocked final CSV export should create a blocked export log entry.")
                assert_true(blocked_log["status"] == "blocked", "Blocked final export should create a blocked export log entry.")
                assert_true(blocked_log["filename"] == "mouse_records_filtered.csv", "Blocked export log should preserve intended filename.")
                export_preview = client.get("/api/export-preview").json()
                assert_true(export_preview["source_layer"] == "export or view", "Export preview should stay an export/view layer.")
                assert_true(export_preview["export_type"] == "separation_preview", "Export preview should identify its workbook-like shape.")
                assert_true(export_preview["preview_row_count"] >= 3, "Export preview should include mouse candidate rows.")
                assert_true(
                    export_preview["latest_data_change_at"] and export_preview["latest_generated_export_at"],
                    "Export preview should expose data and generated export timestamps.",
                )
                assert_true(export_preview["export_stale"] is False, "Freshly generated CSV export should not be stale before more data changes.")
                assert_true(
                    export_preview["expected_separation_filename"].endswith("분리 현황표.xlsx"),
                    "Export preview should expose the expected separation workbook filename.",
                )
                assert_true(
                    export_preview["expected_animal_sheet_filename"].endswith("animal sheet.xlsx"),
                    "Export preview should expose the expected animal sheet workbook filename.",
                )
                assert_true(
                    export_preview["separation_columns"][:5] == ["Cage number", "Strain", "Genotype", "total", "DOB"],
                    "Separation preview should expose senior-workbook-style column labels.",
                )
                assert_true(
                    any(row["total"].endswith("p") and row["source_note_item_ids"] for row in export_preview["separation_rows"]),
                    "Separation preview should group mouse records into workbook-like sex/count rows with traceability.",
                )
                assert_true(
                    any(row["display_id"] == "MT321" and row["source_note_item_id"] for row in export_preview["preview_rows"]),
                    "Export preview rows should preserve source note traceability.",
                )
                mt321_export_row = next(row for row in export_preview["preview_rows"] if row["display_id"] == "MT321")
                assert_true(
                    mt321_export_row["raw_note_line"] == "MT321 R'"
                    and (mt321_export_row["source_photo_id"] or mt321_export_row["card_snapshot_id"]),
                    "Export preview rows should expose raw note and photo/snapshot evidence.",
                )
                assert_true(
                    any(
                        row["display_id"] == "MT399"
                        and row["ear_label_review_status"] == "check"
                        and "Ear label R0: check" in row["uncertainty"]
                        for row in export_preview["preview_rows"]
                    ),
                    "Export preview should surface uncertain ear labels without accepting their normalized code.",
                )
                assert_true(
                    any(row["raw_note_lines"] and (row["source_photo_ids"] or row["card_snapshot_ids"]) for row in export_preview["separation_rows"]),
                    "Separation preview rows should keep raw note and photo/snapshot evidence for traceability.",
                )
                assert_true(
                    export_preview["blocked_review_items"] >= len(export_preview["review_blockers"]),
                    "Export preview should expose review blocker details up to its display limit.",
                )
                assert_true(
                    export_preview["genotype_blocker_items"] >= 1
                    and any(item.get("blocker_type") == "genotype_status" for item in export_preview["readiness_warnings"]),
                    "Export preview should include genotype status readiness warnings from the vocabulary.",
                )
                assert_true(export_preview["ready"] is False, "Open review blockers should keep export preview blocked.")
                dashboard_before = client.get("/api/genotyping-dashboard")
                assert_true(dashboard_before.status_code == 200, "Genotyping dashboard endpoint failed.")
                dashboard_before_rows = {card["key"]: card["count"] for card in dashboard_before.json()}
                assert_true(
                    dashboard_before_rows.get("not_sampled", 0) >= 1,
                    "Genotyping dashboard should count newly separated mice that need sampling.",
                )
                genotyping_target = next(
                    mouse for mouse in mice
                    if mouse["display_id"] == "MT321" and mouse["status"] == "active"
                )
                audit_trace = client.get(f"/api/mice/{genotyping_target['mouse_id']}/audit-trace")
                assert_true(audit_trace.status_code == 200, "Mouse audit trace endpoint failed.")
                audit_payload = audit_trace.json()
                assert_true(
                    audit_payload["source_layer"] == "export or view"
                    and audit_payload["mouse"]["mouse_id"] == genotyping_target["mouse_id"],
                    "Mouse audit trace should identify the selected mouse and boundary.",
                )
                assert_true(
                    any(item["category"] in {"note_line", "mouse_event", "review"} for item in audit_payload["timeline"]),
                    "Mouse audit trace should combine note, event, or review evidence in one timeline.",
                )
                genotyping_evidence_record_id = genotyping_target.get("source_record_id") or (
                    audit_payload["source_records"][0]["source_record_id"] if audit_payload["source_records"] else ""
                )
                assert_true(
                    bool(genotyping_evidence_record_id),
                    "Genotyping result verification needs source evidence for the target mouse.",
                )
                target_rule = client.post(
                    "/api/strain-target-genotypes",
                    json={
                        "strain_text": genotyping_target["raw_strain_text"],
                        "target_genotype": "Tg/Tg",
                        "purpose": "mating_candidate",
                    },
                )
                assert_true(target_rule.status_code == 200, "Could not create configurable target genotype rule.")
                target_rules = client.get("/api/strain-target-genotypes").json()
                assert_true(
                    any(
                        rule["strain_text"] == genotyping_target["raw_strain_text"]
                        and rule["target_genotype"] == "Tg/Tg"
                        and rule["purpose"] == "mating_candidate"
                        for rule in target_rules
                    ),
                    "Target genotype rule API should preserve configured rule values.",
                )
                cage = client.post(
                    "/api/cages",
                    json={
                        "cage_label": "A-101",
                        "location": "Room A",
                        "rack": "R1",
                        "shelf": "S2",
                        "cage_type": "holding",
                        "note": "Verification cage.",
                    },
                )
                assert_true(cage.status_code == 200, "Could not create cage registry entry.")
                cage_payload = cage.json()
                assert_true(cage_payload["source_record_id"], "Cage creation should leave source evidence.")
                duplicate_cage = client.post("/api/cages", json={"cage_label": "a-101"})
                assert_true(duplicate_cage.status_code == 409, "Duplicate cage labels should be rejected case-insensitively.")
                cage_move = client.post(
                    f"/api/mice/{genotyping_target['mouse_id']}/move-cage",
                    json={"cage_id": cage_payload["cage_id"], "note": "Moved for verification."},
                )
                assert_true(cage_move.status_code == 200, "Could not assign mouse to cage.")
                cages = client.get("/api/cages").json()
                assert_true(
                    any(item["cage_label"] == "A-101" and item["active_mouse_count"] == 1 for item in cages),
                    "Cage list should include active assignment counts.",
                )
                female_parent = next(
                    mouse for mouse in mice
                    if mouse["display_id"] == "MT322" and mouse["status"] == "active"
                )
                mating = client.post(
                    "/api/matings",
                    json={
                        "mating_label": "MT321 x MT322",
                        "male_mouse_id": genotyping_target["mouse_id"],
                        "female_mouse_id": female_parent["mouse_id"],
                        "strain_goal": genotyping_target["raw_strain_text"],
                        "expected_genotype": "Tg/Tg",
                        "start_date": "2026-05-01",
                        "purpose": "verification",
                    },
                )
                assert_true(mating.status_code == 200, "Could not create mating registry entry.")
                mating_payload = mating.json()
                mating_rows = client.get("/api/matings").json()
                assert_true(
                    any(
                        row["mating_id"] == mating_payload["mating_id"]
                        and "MT321" in row["male_mice"]
                        and "MT322" in row["female_mice"]
                        for row in mating_rows
                    ),
                    "Mating list should expose linked parent mice by role.",
                )
                missing_parent_mating = client.post(
                    "/api/matings",
                    json={"mating_label": "bad mating", "male_mouse_id": "missing_mouse"},
                )
                assert_true(missing_parent_mating.status_code == 404, "Mating creation should reject missing parent mouse IDs.")
                litter = client.post(
                    "/api/litters",
                    json={
                        "litter_label": "L-MT321-001",
                        "mating_id": mating_payload["mating_id"],
                        "birth_date": "2026-05-02",
                        "number_born": 10,
                        "number_alive": 10,
                        "status": "born",
                    },
                )
                assert_true(litter.status_code == 200, "Could not create litter registry entry.")
                litter_payload = litter.json()
                litter_rows = client.get("/api/litters").json()
                assert_true(
                    any(row["litter_id"] == litter_payload["litter_id"] and row["number_born"] == 10 for row in litter_rows),
                    "Litter list should expose source-backed litter counts.",
                )
                offspring = client.post(
                    f"/api/litters/{litter_payload['litter_id']}/offspring",
                    json={
                        "count": 2,
                        "display_prefix": "MT321-L1",
                        "start_number": 1,
                        "sex": "unknown",
                        "cage_id": cage_payload["cage_id"],
                        "note": "Generated from reviewed litter count.",
                    },
                )
                assert_true(offspring.status_code == 200, "Could not create offspring mouse records from litter.")
                offspring_payload = offspring.json()
                assert_true(offspring_payload["created_count"] == 2, "Offspring creation should return created count.")
                assert_true(offspring_payload["source_record_id"], "Offspring creation should preserve source evidence.")
                duplicate_offspring = client.post(
                    f"/api/litters/{litter_payload['litter_id']}/offspring",
                    json={"count": 1, "display_prefix": "MT321-L1", "start_number": 1},
                )
                assert_true(duplicate_offspring.status_code == 409, "Duplicate offspring IDs should be rejected.")
                offspring_rows = client.get("/api/mice", params={"query": "MT321-L1"}).json()
                assert_true(len(offspring_rows) == 2, "Mouse search should include generated offspring records.")
                assert_true(
                    all(
                        row["father_id"] == genotyping_target["mouse_id"]
                        and row["mother_id"] == female_parent["mouse_id"]
                        and row["litter_id"] == litter_payload["litter_id"]
                        and row["source_record_id"] == offspring_payload["source_record_id"]
                        and row["current_cage_label"] == "A-101"
                        for row in offspring_rows
                    ),
                    "Generated offspring should preserve parent, litter, source, and cage traceability.",
                )
                litter_rows_after_offspring = client.get("/api/litters").json()
                assert_true(
                    any(row["litter_id"] == litter_payload["litter_id"] and row["offspring_count"] == 2 for row in litter_rows_after_offspring),
                    "Litter list should expose generated offspring count.",
                )
                animal_preview = client.get("/api/export-preview").json()
                assert_true(
                    animal_preview["animal_sheet_columns"][:8]
                    == ["Cage No.", "Strain", "Sex", "I.D", "genotype", "DOB", "Mating date", "Pubs"],
                    "Animal sheet preview should expose mating-workbook-style column labels.",
                )
                assert_true(
                    any(
                        row["cage_no"] == "1"
                        and row["strain"] == genotyping_target["raw_strain_text"]
                        and row["mating_date"] == "2026-05-01"
                        for row in animal_preview["animal_sheet_rows"]
                    ),
                    "Animal sheet preview should include parent rows grouped by mating cage.",
                )
                assert_true(
                    any(row["sex"] == "F1" and row["mouse_id"] == "10p" and row["status"] == "pre_weaning" for row in animal_preview["animal_sheet_rows"]),
                    "Animal sheet preview should include litter rows with pup counts and status.",
                )
                over_weaned = client.post(
                    f"/api/litters/{litter_payload['litter_id']}/wean",
                    json={"weaning_date": "2026-05-23", "number_weaned": 3},
                )
                assert_true(over_weaned.status_code == 409, "Weaning should reject counts above generated offspring records.")
                weaned = client.post(
                    f"/api/litters/{litter_payload['litter_id']}/wean",
                    json={
                        "weaning_date": "2026-05-23",
                        "number_weaned": 2,
                        "note": "Verified weaning from reviewed litter card.",
                    },
                )
                assert_true(weaned.status_code == 200, "Could not complete litter weaning.")
                weaned_payload = weaned.json()
                assert_true(weaned_payload["status"] == "weaned", "Weaning should mark litter status as weaned.")
                assert_true(weaned_payload["number_weaned"] == 2, "Weaning should preserve reviewed weaned count.")
                assert_true(weaned_payload["source_record_id"], "Weaning should preserve source evidence.")
                duplicate_wean = client.post(
                    f"/api/litters/{litter_payload['litter_id']}/wean",
                    json={"weaning_date": "2026-05-24", "number_weaned": 2},
                )
                assert_true(duplicate_wean.status_code == 409, "Already-weaned litters should not be silently overwritten.")
                weaned_offspring_rows = client.get("/api/mice", params={"query": "MT321-L1"}).json()
                assert_true(
                    all(row["status"] == "active" and row["next_action"] == "sample_needed" for row in weaned_offspring_rows),
                    "Weaned offspring should move from weaning pending to active sample-needed workflow.",
                )
                requested_genotyping = client.post(
                    "/api/genotyping/request",
                    json={
                        "mouse_id": weaned_offspring_rows[0]["mouse_id"],
                        "sample_id": "TAIL-MT321-L1-01",
                        "target_name": "ApoM Tg/Tg",
                        "note": "Requested after weaning verification.",
                    },
                )
                assert_true(requested_genotyping.status_code == 200, "Could not request genotyping for weaned offspring.")
                requested_payload = requested_genotyping.json()
                assert_true(requested_payload["sample_id"] == "TAIL-MT321-L1-01", "Genotyping request should preserve sample ID.")
                assert_true(requested_payload["genotyping_status"] == "submitted", "Genotyping request should mark mouse submitted.")
                assert_true(requested_payload["next_action"] == "awaiting_result", "Genotyping request should move mouse to awaiting result.")
                litter_rows_after_weaning = client.get("/api/litters").json()
                assert_true(
                    any(
                        row["litter_id"] == litter_payload["litter_id"]
                        and row["status"] == "weaned"
                        and row["number_weaned"] == 2
                        and row["weaning_date"] == "2026-05-23"
                        for row in litter_rows_after_weaning
                    ),
                    "Litter list should expose completed weaning state.",
                )
                genotyping_update = client.post(
                    "/api/genotyping/update",
                    json={
                        "mouse_id": genotyping_target["mouse_id"],
                        "sample_id": "MT321",
                        "raw_result": "Tg/Tg",
                        "normalized_result": "Tg/Tg",
                        "source_record_id": genotyping_evidence_record_id,
                    },
                )
                assert_true(
                    genotyping_update.status_code == 200,
                    f"Could not update genotyping workflow state: {genotyping_update.status_code} {genotyping_update.text}",
                )
                genotyping_payload = genotyping_update.json()
                assert_true(genotyping_payload["genotyping_status"] == "resulted", "Genotyping result should mark mouse resulted.")
                assert_true(genotyping_payload["target_match_status"] == "matches_target", "Genotyping result should use configured target matching.")
                assert_true(genotyping_payload["next_action"] == "consider_for_mating", "Matching target genotype should suggest a mating review action.")
                experiment_readiness = client.get("/api/experiment-readiness").json()
                ready_experiment_mouse = next(
                    item for item in experiment_readiness["rows"]
                    if item["mouse_id"] == genotyping_target["mouse_id"]
                )
                assert_true(
                    ready_experiment_mouse["readiness_status"] == "warning"
                    and ready_experiment_mouse["genotype_status"] == "confirmed",
                    "Experiment readiness should show confirmed non-experimental-use mice as warnings, not genotype blockers.",
                )
                stale_preview = client.get("/api/export-preview").json()
                assert_true(
                    stale_preview["export_stale"] is True
                    and stale_preview["latest_data_change_at"] >= stale_preview["latest_generated_export_at"],
                    "Export preview should warn when mouse data changed after the last generated export.",
                )
                duplicate_resulted_request = client.post(
                    "/api/genotyping/request",
                    json={"mouse_id": genotyping_target["mouse_id"], "sample_id": "already-resulted"},
                )
                assert_true(duplicate_resulted_request.status_code == 409, "Resulted mice should not accept a new genotyping request silently.")
                parent_trace = client.get(f"/api/mice/{genotyping_target['mouse_id']}/audit-trace")
                assert_true(parent_trace.status_code == 200, "Mouse audit trace endpoint should return parent mouse evidence.")
                parent_trace_payload = parent_trace.json()
                parent_categories = {item["category"] for item in parent_trace_payload["timeline"]}
                assert_true(parent_trace_payload["source_layer"] == "export or view", "Mouse audit trace should stay a review/export view.")
                assert_true(parent_trace_payload["note_items"], "Mouse audit trace should include parsed note-line evidence.")
                assert_true("note_line" in parent_categories, "Mouse audit trace timeline should include note-line evidence.")
                assert_true("genotyping" in parent_categories, "Mouse audit trace timeline should include genotyping records.")
                offspring_trace = client.get(f"/api/mice/{weaned_offspring_rows[0]['mouse_id']}/audit-trace")
                assert_true(offspring_trace.status_code == 200, "Mouse audit trace endpoint should return offspring evidence.")
                offspring_trace_payload = offspring_trace.json()
                offspring_categories = {item["category"] for item in offspring_trace_payload["timeline"]}
                assert_true("mouse_event" in offspring_categories, "Offspring audit trace should include born/weaned/request events.")
                assert_true("genotyping" in offspring_categories, "Offspring audit trace should include pending genotyping request.")
                assert_true(
                    offspring_trace_payload["source_records"],
                    "Offspring audit trace should preserve source records linked through mouse events.",
                )
                dashboard_after = {card["key"]: card["count"] for card in client.get("/api/genotyping-dashboard").json()}
                assert_true(
                    dashboard_after.get("target_confirmed", 0) >= 1,
                    "Genotyping dashboard should count mice with confirmed target genotypes.",
                )
                genotyping_records = client.get("/api/genotyping-records").json()
                assert_true(
                    any(record["mouse_id"] == genotyping_target["mouse_id"] and record["normalized_result"] == "Tg/Tg" for record in genotyping_records),
                    "Genotyping record history should preserve the entered result.",
                )
                assert_true(
                    any(record["sample_id"] == "TAIL-MT321-L1-01" and record["result_status"] == "pending" for record in genotyping_records),
                    "Genotyping request should create a pending genotyping record.",
                )
                missing_audit = client.get("/api/mice/mouse-does-not-exist/audit-trail")
                assert_true(missing_audit.status_code == 404, "Missing mouse audit trail should return 404.")
                audit_trace = client.get(f"/api/mice/{genotyping_target['mouse_id']}/audit-trail")
                assert_true(audit_trace.status_code == 200, "Mouse audit trail endpoint failed.")
                audit_payload = audit_trace.json()
                assert_true(audit_payload["source_layer"] == "export or view", "Audit trail should be a read-only export/view layer.")
                assert_true(audit_payload["mouse"]["display_id"] == "MT321", "Audit trail should include the selected mouse.")
                audit_categories = {item["category"] for item in audit_payload["timeline"]}
                assert_true(
                    {"note_line", "mouse_event", "genotyping", "cage_assignment", "action_log"}.issubset(audit_categories),
                    "Audit trail should combine source notes, events, genotyping, cage history, and action logs.",
                )
                assert_true(
                    any(action["action_type"] == "genotyping_resulted" for action in audit_payload["actions"]),
                    "Audit trail should expose direct action log records for the mouse.",
                )
                assert_true(
                    any(assignment["cage_label"] == "A-101" for assignment in audit_payload["cage_assignments"]),
                    "Audit trail should expose cage assignment history.",
                )
                offspring_audit = client.get(f"/api/mice/{requested_payload['mouse_id']}/audit-trail").json()
                assert_true(
                    offspring_audit["lineage"]["litter"]["litter_id"] == litter_payload["litter_id"]
                    and offspring_audit["lineage"]["father"]["mouse_id"] == genotyping_target["mouse_id"],
                    "Offspring audit trail should expose litter and parent lineage.",
                )
                genotyping_export = client.get("/api/exports/genotyping-worklist.csv", params={"query": "MT321"})
                assert_true(genotyping_export.status_code == 200, "Genotyping worklist CSV export endpoint failed.")
                assert_true(
                    "genotyping_worklist_filtered.csv" in genotyping_export.headers.get("content-disposition", ""),
                    "Filtered genotyping worklist should use the filtered filename.",
                )
                assert_true(
                    "target_match_status" in genotyping_export.text
                    and "matches_target" in genotyping_export.text
                    and "consider_for_mating" in genotyping_export.text,
                    "Genotyping worklist export should include target match and next action fields.",
                )
                genotyping_export_log = client.get("/api/export-log").json()
                assert_true(
                    genotyping_export_log[0]["export_type"] == "genotyping_worklist_csv",
                    "Export log should record companion genotyping worklist exports.",
                )
                filtered_mice = client.get("/api/mice", params={"query": "MT321"}).json()
                assert_true(
                    filtered_mice and all("MT321" in mouse["display_id"] for mouse in filtered_mice),
                    "Mouse API query should filter mouse records.",
                )
                search = client.get("/api/search", params={"query": "Tg/Tg"}).json()
                assert_true(
                    any(mouse["display_id"] == "MT321" for mouse in search["mice"]),
                    "Search API should include matching genotyping mouse records.",
                )
                csv_export = client.get("/api/exports/mice.csv", params={"query": "MT321"})
                assert_true(csv_export.status_code == 200, "Mouse CSV export endpoint failed.")
                assert_true(
                    "mouse_records_filtered.csv" in csv_export.headers.get("content-disposition", ""),
                    "Filtered CSV export should use the filtered filename.",
                )
                assert_true(
                    "display_id" in csv_export.text and "MT321" in csv_export.text,
                    "Mouse CSV export should include headers and filtered mouse rows.",
                )
                offspring_csv = client.get("/api/exports/mice.csv", params={"query": "MT321-L1"})
                assert_true(offspring_csv.status_code == 200, "Offspring CSV export endpoint failed.")
                assert_true(
                    "father_id" in offspring_csv.text
                    and "mother_id" in offspring_csv.text
                    and "litter_id" in offspring_csv.text
                    and offspring_payload["source_record_id"] in offspring_csv.text,
                    "Offspring CSV export should include lineage and source traceability fields.",
                )
                with db.connection() as conn:
                    note_count = conn.execute(
                        "SELECT COUNT(*) AS count FROM card_note_item_log"
                    ).fetchone()["count"]
                    mouse_count = conn.execute("SELECT COUNT(*) AS count FROM mouse_master").fetchone()["count"]
                    moved_count = conn.execute(
                        "SELECT COUNT(*) AS count FROM mouse_master WHERE status = 'moved'"
                    ).fetchone()["count"]
                    duplicate_leak_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM mouse_master
                        WHERE source_note_item_id LIKE ?
                        """,
                        ("note_FIXTURE-DUPLICATE-ACTIVE_%",),
                    ).fetchone()["count"]
                    genotyping_action_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM action_log
                        WHERE action_type = 'genotyping_resulted' AND target_id = ?
                        """,
                        (genotyping_target["mouse_id"],),
                    ).fetchone()["count"]
                    cage_move_action_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM action_log
                        WHERE action_type = 'mouse_cage_moved' AND target_id = ?
                        """,
                        (genotyping_target["mouse_id"],),
                    ).fetchone()["count"]
                    cage_move_event_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM mouse_event
                        WHERE event_type = 'moved' AND related_entity_type = 'cage'
                        """
                    ).fetchone()["count"]
                    active_cage_assignment_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM mouse_cage_assignment
                        WHERE mouse_id = ? AND status = 'active'
                        """,
                        (genotyping_target["mouse_id"],),
                    ).fetchone()["count"]
                    ended_cage_assignment_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM mouse_cage_assignment
                        WHERE mouse_id = ? AND status = 'ended'
                        """,
                        (genotyping_target["mouse_id"],),
                    ).fetchone()["count"]
                    mating_action_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM action_log
                        WHERE action_type = 'mating_created' AND target_id = ?
                        """,
                        (mating_payload["mating_id"],),
                    ).fetchone()["count"]
                    litter_action_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM action_log
                        WHERE action_type = 'litter_created' AND target_id = ?
                        """,
                        (litter_payload["litter_id"],),
                    ).fetchone()["count"]
                    paired_event_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM mouse_event
                        WHERE event_type = 'paired' AND related_entity_type = 'mating'
                        """
                    ).fetchone()["count"]
                    litter_event_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM mouse_event
                        WHERE event_type = 'litter_produced' AND related_entity_type = 'litter'
                        """
                    ).fetchone()["count"]
                    offspring_born_event_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM mouse_event
                        WHERE event_type = 'born' AND related_entity_type = 'litter'
                        """
                    ).fetchone()["count"]
                    offspring_action_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM action_log
                        WHERE action_type = 'offspring_created' AND target_id = ?
                        """,
                        (litter_payload["litter_id"],),
                    ).fetchone()["count"]
                    weaned_event_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM mouse_event
                        WHERE event_type = 'weaned' AND related_entity_type = 'litter'
                        """
                    ).fetchone()["count"]
                    weaned_action_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM action_log
                        WHERE action_type = 'litter_weaned' AND target_id = ?
                        """,
                        (litter_payload["litter_id"],),
                    ).fetchone()["count"]
                    genotyping_request_action_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM action_log
                        WHERE action_type = 'genotyping_requested'
                        """
                    ).fetchone()["count"]
                    tail_biopsy_event_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM mouse_event
                        WHERE event_type = 'tail_biopsy' AND related_entity_type = 'genotyping_record'
                        """
                    ).fetchone()["count"]
                    genotyping_request_event_count = conn.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM mouse_event
                        WHERE event_type = 'genotyping_requested' AND related_entity_type = 'genotyping_record'
                        """
                    ).fetchone()["count"]
                assert_true(note_count >= 10, "Persisted note item evidence count is too low.")
                assert_true(mouse_count >= 5, "Persisted mouse candidate count is too low.")
                assert_true(moved_count >= 1, "Single-struck mouse note should create a moved candidate.")
                assert_true(duplicate_leak_count == 0, "Duplicate active fixture should not create mouse candidates.")
                assert_true(genotyping_action_count == 1, "Genotyping update should create an action log entry.")
                assert_true(cage_move_action_count >= 2, "Each cage move should create an action log entry.")
                assert_true(cage_move_event_count >= 1, "Cage move should create a mouse event.")
                assert_true(active_cage_assignment_count == 1, "Cage moves should leave only one active assignment per mouse.")
                assert_true(ended_cage_assignment_count >= 1, "Cage moves should close the previous active assignment.")
                assert_true(mating_action_count == 1, "Mating creation should create an action log entry.")
                assert_true(litter_action_count == 1, "Litter creation should create an action log entry.")
                assert_true(paired_event_count >= 2, "Mating creation should create parent pairing events.")
                assert_true(litter_event_count >= 2, "Litter creation should create parent litter events.")
                assert_true(offspring_born_event_count == 2, "Offspring creation should create one birth event per mouse.")
                assert_true(offspring_action_count == 1, "Offspring creation should create an action log entry.")
                assert_true(weaned_event_count == 2, "Weaning should create one event per weaned offspring.")
                assert_true(weaned_action_count == 1, "Weaning should create an action log entry.")
                assert_true(genotyping_request_action_count == 1, "Genotyping request should create an action log entry.")
                assert_true(tail_biopsy_event_count == 1, "Genotyping request should create a tail biopsy event.")
                assert_true(genotyping_request_event_count == 1, "Genotyping request should create a request event.")
                correction = client.post(
                    "/api/corrections",
                    json={
                        "entity_type": "strain",
                        "entity_id": strain_payload["strain_id"],
                        "field_name": "common_name",
                        "before_value": "",
                        "after_value": "PV-Cre line",
                        "reason": "Verified local correction workflow.",
                        "source_record_id": strain_payload["source_record_id"],
                    },
                )
                assert_true(correction.status_code == 200, "Could not create correction log entry.")
                corrections = client.get("/api/corrections").json()
                assert_true(
                    corrections[0]["before_value"] == "" and corrections[0]["after_value"] == "PV-Cre line",
                    "Correction log should preserve before and after values.",
                )
                remaining_reviews = [
                    item for item in client.get("/api/review-items").json() if item["status"] == "open"
                ]
                assert_true(remaining_reviews, "Fixture should still have open review blockers before final release.")
                for item in remaining_reviews:
                    release_review = client.post(
                        f"/api/review-items/{item['review_id']}/resolve",
                        json={
                            "resolution_note": "Verified blocker before releasing ready CSV export.",
                            "resolved_value": item.get("suggested_value") or item.get("current_value") or "",
                        },
                    )
                    assert_true(
                        release_review.status_code == 200,
                        f"Could not resolve remaining review blocker: {release_review.status_code} {release_review.text}",
                    )
                ready_preview = client.get("/api/export-preview").json()
                assert_true(ready_preview["ready"] is True, "Resolved review blockers should make export preview ready.")
                assert_true(ready_preview["blocked_review_items"] == 0, "Ready export preview should have no blockers.")
                assert_true(
                    "readiness_warnings" in ready_preview
                    and "genotype_blocker_items" in ready_preview,
                    "Ready export preview should still expose genotype readiness warnings separately from export blockers.",
                )
                ready_export = client.get("/api/exports/mice.csv", params={"query": "MT321", "require_ready": "true"})
                assert_true(ready_export.status_code == 200, "Ready CSV export should succeed after review resolution.")
                assert_true("MT321" in ready_export.text, "Ready CSV export should include the filtered mouse row.")
                post_ready_preview = client.get("/api/export-preview").json()
                assert_true(post_ready_preview["export_stale"] is False, "Generated ready export should clear the stale export warning.")
                ready_separation_xlsx = client.get("/api/exports/separation.xlsx")
                assert_true(ready_separation_xlsx.status_code == 200, "Ready separation XLSX export should succeed after review resolution.")
                assert_true(
                    ready_separation_xlsx.content[:4] == b"PK\x03\x04",
                    "Separation XLSX export should be a ZIP-based workbook.",
                )
                separation_disposition = ready_separation_xlsx.headers.get("content-disposition", "")
                assert_true(
                    "filename*=UTF-8''" in separation_disposition and "separation.xlsx" in separation_disposition,
                    "Separation XLSX export should expose a safe fallback and UTF-8 filename.",
                )
                with zipfile.ZipFile(io.BytesIO(ready_separation_xlsx.content)) as workbook_zip:
                    assert_true(
                        "xl/workbook.xml" in workbook_zip.namelist()
                        and "xl/styles.xml" in workbook_zip.namelist()
                        and "xl/worksheets/sheet1.xml" in workbook_zip.namelist()
                        and "xl/worksheets/sheet2.xml" in workbook_zip.namelist(),
                        "Separation XLSX should contain the required workbook parts.",
                    )
                    separation_workbook_xml = workbook_zip.read("xl/workbook.xml").decode("utf-8")
                    separation_sheet_xml = workbook_zip.read("xl/worksheets/sheet1.xml").decode("utf-8")
                    separation_trace_xml = workbook_zip.read("xl/worksheets/sheet2.xml").decode("utf-8")
                    separation_styles_xml = workbook_zip.read("xl/styles.xml").decode("utf-8")
                assert_true("분리 현황표" in separation_workbook_xml, "Separation XLSX should use the lab workbook sheet name.")
                assert_true("Sampling point" in separation_sheet_xml, "Separation XLSX should include the template header.")
                assert_true("ApoM Tg/Tg" in separation_sheet_xml, "Separation XLSX should include accepted strain rows.")
                assert_true("<cols>" in separation_sheet_xml and 's="1"' in separation_sheet_xml, "Separation XLSX should include column widths and styled headers.")
                assert_true("Export_Trace" in separation_workbook_xml and "Source note" in separation_trace_xml, "Separation XLSX should include traceability sheet.")
                assert_true("<b/>" in separation_styles_xml, "Separation XLSX should include bold header style.")
                assert_true(load_workbook is not None, "openpyxl is required to validate generated XLSX workbooks.")
                separation_workbook = load_workbook(io.BytesIO(ready_separation_xlsx.content), data_only=True)
                assert_true("분리 현황표" in separation_workbook.sheetnames, "openpyxl should load the separation sheet name.")
                assert_true("Export_Trace" in separation_workbook.sheetnames, "openpyxl should load the separation trace sheet.")
                separation_sheet = separation_workbook["분리 현황표"]
                assert_true(
                    [separation_sheet.cell(1, column).value for column in range(1, 9)]
                    == ["Cage number", "Strain", "Genotype", "total", "DOB", "WT", "Tg", "Sampling point"],
                    "openpyxl should read the separation workbook headers.",
                )
                assert_true(
                    any(row[1] == "ApoM Tg/Tg" for row in separation_sheet.iter_rows(min_row=2, values_only=True)),
                    "openpyxl should read accepted separation strain rows.",
                )
                separation_trace_sheet = separation_workbook["Export_Trace"]
                assert_true(
                    separation_trace_sheet.cell(1, 2).value == "Source note"
                    and separation_trace_sheet.cell(1, 4).value == "Boundary",
                    "openpyxl should read separation trace headers.",
                )
                assert_true(
                    separation_trace_sheet.cell(1, 6).value == "Source photo"
                    and separation_trace_sheet.cell(1, 8).value == "Raw note line"
                    and separation_trace_sheet.cell(1, 9).value == "Uncertainty",
                    "Separation trace sheet should include photo, raw note, and uncertainty columns.",
                )
                ready_animal_xlsx = client.get("/api/exports/animal-sheet.xlsx")
                assert_true(ready_animal_xlsx.status_code == 200, "Ready animal sheet XLSX export should succeed after review resolution.")
                assert_true(
                    ready_animal_xlsx.content[:4] == b"PK\x03\x04",
                    "Animal sheet XLSX export should be a ZIP-based workbook.",
                )
                animal_disposition = ready_animal_xlsx.headers.get("content-disposition", "")
                assert_true(
                    "filename*=UTF-8''" in animal_disposition and "animal sheet.xlsx" in animal_disposition,
                    "Animal sheet XLSX export should expose a safe fallback and UTF-8 filename.",
                )
                with zipfile.ZipFile(io.BytesIO(ready_animal_xlsx.content)) as workbook_zip:
                    assert_true(
                        "xl/workbook.xml" in workbook_zip.namelist()
                        and "xl/styles.xml" in workbook_zip.namelist()
                        and "xl/worksheets/sheet1.xml" in workbook_zip.namelist()
                        and "xl/worksheets/sheet2.xml" in workbook_zip.namelist(),
                        "Animal sheet XLSX should contain the required workbook parts.",
                    )
                    animal_workbook_xml = workbook_zip.read("xl/workbook.xml").decode("utf-8")
                    animal_sheet_xml = workbook_zip.read("xl/worksheets/sheet1.xml").decode("utf-8")
                    animal_trace_xml = workbook_zip.read("xl/worksheets/sheet2.xml").decode("utf-8")
                assert_true("Cage No." in animal_sheet_xml, "Animal sheet XLSX should include the template header.")
                assert_true("MT321" in animal_sheet_xml and "MT322" in animal_sheet_xml, "Animal sheet XLSX should include parent IDs.")
                assert_true("2026-05-02 10p" in animal_sheet_xml, "Animal sheet XLSX should include litter pup evidence.")
                assert_true("animal sheet" in animal_workbook_xml and "Export_Trace" in animal_workbook_xml, "Animal sheet XLSX should name workbook sheets clearly.")
                assert_true("<cols>" in animal_sheet_xml and 's="1"' in animal_sheet_xml, "Animal sheet XLSX should include column widths and styled headers.")
                assert_true("Source record" in animal_trace_xml, "Animal sheet XLSX should include traceability sheet.")
                animal_workbook = load_workbook(io.BytesIO(ready_animal_xlsx.content), data_only=True)
                assert_true("animal sheet" in animal_workbook.sheetnames, "openpyxl should load the animal sheet name.")
                assert_true("Export_Trace" in animal_workbook.sheetnames, "openpyxl should load the animal trace sheet.")
                animal_sheet = animal_workbook["animal sheet"]
                assert_true(
                    [animal_sheet.cell(1, column).value for column in range(1, 9)]
                    == ["Cage No.", "Strain", "Sex", "I.D", "genotype", "DOB", "Mating date", "Pubs"],
                    "openpyxl should read the animal sheet headers.",
                )
                animal_values = list(animal_sheet.iter_rows(min_row=2, values_only=True))
                assert_true(
                    any("MT321" in str(cell) for row in animal_values for cell in row if cell is not None)
                    and any("MT322" in str(cell) for row in animal_values for cell in row if cell is not None),
                    "openpyxl should read animal sheet parent IDs.",
                )
                assert_true(
                    any("2026-05-02 10p" in str(cell) for row in animal_values for cell in row if cell is not None),
                    "openpyxl should read animal sheet litter pup evidence.",
                )
                animal_trace_sheet = animal_workbook["Export_Trace"]
                assert_true(
                    animal_trace_sheet.cell(1, 3).value == "Source record"
                    and animal_trace_sheet.cell(1, 4).value == "Boundary",
                    "openpyxl should read animal trace headers.",
                )
                assert_true(
                    animal_trace_sheet.cell(1, 6).value == "Source photo"
                    and animal_trace_sheet.cell(1, 8).value == "Raw note line"
                    and animal_trace_sheet.cell(1, 9).value == "Uncertainty",
                    "Animal trace sheet should include photo, raw note, and uncertainty columns.",
                )
                ready_logs = client.get("/api/export-log").json()
                assert_true(
                    any(item["export_type"] == "separation_xlsx" and item["status"] == "generated" for item in ready_logs),
                    "Export log should record generated separation XLSX exports.",
                )
                assert_true(
                    any(item["export_type"] == "animal_sheet_xlsx" and item["status"] == "generated" for item in ready_logs),
                    "Export log should record generated animal sheet XLSX exports.",
                )
                workbook_export_log = next(
                    item for item in ready_logs if item["export_type"] == "separation_xlsx" and item["status"] == "generated"
                )
                assert_true(
                    workbook_export_log["export_manifest_path"].endswith(".json")
                    and workbook_export_log["validation_report_id"].startswith("validation_report_export_separation_xlsx")
                    and workbook_export_log["state_watermark"],
                    "Workbook export log should expose manifest, validation report, and state watermark provenance.",
                )
                manifest_preview = client.get(
                    "/api/artifacts/preview",
                    params={"path": workbook_export_log["export_manifest_path"]},
                )
                assert_true(manifest_preview.status_code == 200, f"Could not preview export manifest artifact: {manifest_preview.text}")
                manifest_preview_payload = manifest_preview.json()
                assert_true(
                    manifest_preview_payload["artifact_type"] == "export_manifest"
                    and manifest_preview_payload["source_layer"] == "export or view"
                    and manifest_preview_payload["artifact"]["validation_report_id"],
                    "Artifact preview should expose the export manifest JSON without treating it as canonical state.",
                )
                assert_true(
                    workbook_export_log["validation_report_path"].endswith(".json"),
                    "Workbook export log should expose the validation report artifact path from its manifest.",
                )
                report_preview = client.get(
                    "/api/artifacts/preview",
                    params={"path": workbook_export_log["validation_report_path"]},
                )
                assert_true(report_preview.status_code == 200, f"Could not preview export validation report artifact: {report_preview.text}")
                report_preview_payload = report_preview.json()
                assert_true(
                    report_preview_payload["artifact_type"] == "validation_report"
                    and report_preview_payload["source_layer"] == "export or view"
                    and report_preview_payload["artifact"]["scope"] == "export",
                    "Artifact preview should expose the export validation report JSON as a review artifact.",
                )
                ready_export_log = next(item for item in ready_logs if item["export_type"] == "mouse_csv")
                assert_true(ready_export_log["status"] == "generated", "Ready export should create a generated export log entry.")
                assert_true(
                    ready_export_log["blocked_review_count"] == 0,
                    "Ready export log should record zero review blockers after resolution.",
                )
                void_candidate = client.post(f"/api/canonical-candidates/{mapped_payload['canonical_candidate_id']}/void")
                assert_true(void_candidate.status_code == 200, f"Could not void applied canonical candidate: {void_candidate.text}")
                void_payload = void_candidate.json()
                assert_true(
                    void_payload["boundary"] == "canonical structured state",
                    "Voiding an applied candidate should be classified as canonical structured state.",
                )
                assert_true(
                    void_payload["updated_mice"] == 2 and void_payload["created_events"] == 2,
                    "Voiding should mark applied mouse records without deleting their evidence-backed rows.",
                )
                voided_candidates = client.get("/api/canonical-candidates").json()
                voided_candidate = next(
                    item for item in voided_candidates
                    if item["candidate_id"] == mapped_payload["canonical_candidate_id"]
                )
                assert_true(voided_candidate["status"] == "voided", "Voided canonical candidate should expose voided status.")
                voided_mice = client.get("/api/mice").json()
                voided_mt321 = next((mouse for mouse in voided_mice if mouse["display_id"] == "MT321"), None)
                voided_mt322 = next((mouse for mouse in voided_mice if mouse["display_id"] == "MT322"), None)
                assert_true(
                    voided_mt321 is not None
                    and voided_mt322 is not None
                    and voided_mt321["status"] == "voided"
                    and voided_mt322["status"] == "voided",
                    "Voiding should preserve mouse records and mark them voided instead of deleting them.",
                )
                assert_true(
                    voided_mt321["source_note_item_id"] and voided_mt322["source_note_item_id"],
                    "Voided mouse records should retain source note-line anchors.",
                )
                void_events = [
                    item
                    for item in client.get("/api/mouse-events").json()
                    if item["related_entity_id"] == mapped_payload["canonical_candidate_id"]
                    and item["event_type"] == "canonical_candidate_voided"
                ]
                assert_true(len(void_events) == 2, "Voiding should create explicit candidate void events.")
                voided_audit = client.get(f"/api/canonical-candidates/{mapped_payload['canonical_candidate_id']}/audit").json()
                assert_true(
                    voided_audit["can_void"] is False
                    and voided_audit["summary"]["voided_event_count"] == 2,
                    "Audit should show that an already voided candidate cannot be voided again.",
                )
                idempotent_void = client.post(f"/api/canonical-candidates/{mapped_payload['canonical_candidate_id']}/void")
                assert_true(
                    idempotent_void.status_code == 409,
                    "Re-voiding a voided canonical candidate should be blocked.",
                )
                apply_after_void = client.post(f"/api/canonical-candidates/{mapped_payload['canonical_candidate_id']}/apply")
                assert_true(
                    apply_after_void.status_code == 409,
                    "Applying a voided canonical candidate should be blocked.",
                )

    print("Local app scaffold verification passed.")


if __name__ == "__main__":
    main()
